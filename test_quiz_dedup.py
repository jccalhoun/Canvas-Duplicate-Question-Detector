"""Tests for quiz_dedup.py. Run with:  python3 -m unittest -v

The exports here are synthetic: small QTI files written the way Canvas
writes Classic Quizzes and New Quizzes questions in a course export. They
check the matching logic, not every quirk of a real export."""
import contextlib
import html
import io
import json
import os
import sys
import tempfile
import unittest
import zipfile
from unittest import mock

import xml.etree.ElementTree as ET

import quiz_dedup as qd

NS = "http://www.imsglobal.org/xsd/ims_qtiasiv1p2"
MC = "multiple_choice_question"
_n = [1000]


def item(qtype, stem, choices=None, correct=(), answers=None, rows=None, new=False, bank_ref=None):
    """One <item>. choices/correct for choice questions, answers for short
    answer, rows=[(prompt, answer)] for matching. new=True writes it the way
    a New Quizzes question is exported: uuid answer ids, choices wrapped in
    <p>, an item id without the leading g, and a calculator_type field."""
    _n[0] += 1
    base = _n[0] * 100
    body, cond = "", '<respcondition continue="No"><conditionvar><other/></conditionvar></respcondition>'
    mat = lambda t: f'<material><mattext texttype="text/html">{html.escape(t)}</mattext></material>'
    if rows:
        opts = sorted({a for _, a in rows})
        ids = {a: str(base + k) for k, a in enumerate(opts)}
        labels = "".join(f'<response_label ident="{ids[a]}">{mat(a)}</response_label>' for a in opts)
        body = "".join(f'<response_lid ident="response_{k}">{mat(p)}<render_choice>{labels}</render_choice></response_lid>'
                       for k, (p, _) in enumerate(rows))
        cond = "".join(f'<respcondition><conditionvar><varequal respident="response_{k}">{ids[a]}</varequal></conditionvar>'
                       f'<setvar varname="SCORE" action="Add">25.00</setvar></respcondition>' for k, (_, a) in enumerate(rows))
    elif choices:
        ids = {c: f"{base + k:08x}-0000-4000-8000-000000000000" if new else str(base + k)
               for k, c in enumerate(choices)}
        wrap = (lambda c: f"<p>{c}</p>") if new else (lambda c: c)
        labels = "".join(f'<response_label ident="{ids[c]}">{mat(wrap(c))}</response_label>' for c in choices)
        body = f'<response_lid ident="response1" rcardinality="Single"><render_choice>{labels}</render_choice></response_lid>'
        if len(correct) == 1:
            test = f'<varequal respident="response1">{ids[correct[0]]}</varequal>'
        else:
            test = "<and>" + "".join(
                f'<varequal respident="response1">{ids[c]}</varequal>' if c in correct
                else f'<not><varequal respident="response1">{ids[c]}</varequal></not>' for c in choices) + "</and>"
        cond = f'<respcondition continue="No"><conditionvar>{test}</conditionvar><setvar action="Set" varname="SCORE">100</setvar></respcondition>'
    elif answers:
        body = '<response_str ident="response1" rcardinality="Single"><render_fib><response_label ident="answer1"/></render_fib></response_str>'
        cond = ('<respcondition continue="No"><conditionvar>'
                + "".join(f'<varequal respident="response1">{a}</varequal>' for a in answers)
                + '</conditionvar><setvar action="Set" varname="SCORE">100</setvar></respcondition>')
    ident = f"{_n[0]:032x}" if new else f"g{_n[0]}"
    calculator = ('<qtimetadatafield><fieldlabel>calculator_type</fieldlabel><fieldentry>none</fieldentry>'
                  '</qtimetadatafield>') if new else ""
    if bank_ref:
        # how a Classic quiz question copied from a bank points back at the bank question
        calculator += ('<qtimetadatafield><fieldlabel>assessment_question_identifierref</fieldlabel>'
                       f'<fieldentry>{bank_ref}</fieldentry></qtimetadatafield>')
    return (f'<item ident="{ident}" title="Question"><itemmetadata><qtimetadata>'
            f'<qtimetadatafield><fieldlabel>question_type</fieldlabel><fieldentry>{qtype}</fieldentry></qtimetadatafield>'
            f'<qtimetadatafield><fieldlabel>points_possible</fieldlabel><fieldentry>1.0</fieldentry></qtimetadatafield>'
            f'{calculator}</qtimetadata></itemmetadata><presentation>{mat("<div><p>" + stem + "</p></div>")}{body}</presentation>'
            f'<resprocessing>{cond}</resprocessing></item>')


def item_id(xml):
    """The ident of an <item> built by item()."""
    return xml.split('ident="', 1)[1].split('"', 1)[0]


def quiz(ident, title, items, bank_refs=(), draws=(), new=False, picks=()):
    """A quiz. bank_refs are question groups pointing at a bank; draws are
    the same with a pick count, [(bank id, how many)], where None leaves the
    count empty. picks are single bank questions used by reference,
    [(bank id, item id)]. new=True marks it the way a New Quiz is exported."""
    groups = "".join(f'<section ident="grp{k}" title="Group"><selection_ordering><selection>'
                     f'<sourcebank_ref>{b}</sourcebank_ref></selection></selection_ordering></section>'
                     for k, b in enumerate(bank_refs))
    groups += "".join(f'<section ident="draw{k}" title="Group"><selection_ordering><selection>'
                      f'<sourcebank_ref>{b}</sourcebank_ref>'
                      f'{"<selection_number/>" if n is None else f"<selection_number>{n}</selection_number>"}'
                      f'<selection_extension><points_per_item>1.0</points_per_item></selection_extension>'
                      f'</selection></selection_ordering></section>' for k, (b, n) in enumerate(draws))
    groups += "".join(f'<bankentry_item sourcebank_ref="{b}" item_ref="{i}" points_possible="1.0" entry_type="Item"/>'
                      for b, i in picks)
    marker = ' external_assignment_id="18355351"' if new else ""
    return (f'<?xml version="1.0" encoding="UTF-8"?>\n<questestinterop xmlns="{NS}">'
            f'<assessment ident="{ident}" title="{title}"{marker}><section ident="root_section">'
            f'{"".join(items)}{groups}</section></assessment></questestinterop>')


def bank(ident, title, items, new=False):
    """A question bank, or with new=True a New Quizzes item bank."""
    marker = ' canvas_item_bank="true"' if new else ""
    return (f'<?xml version="1.0" encoding="UTF-8"?>\n<questestinterop xmlns="{NS}"><objectbank ident="{ident}"{marker}>'
            f'<qtimetadata><qtimetadatafield><fieldlabel>bank_title</fieldlabel><fieldentry>{title}</fieldentry>'
            f'</qtimetadatafield></qtimetadata>{"".join(items)}</objectbank></questestinterop>')


class ExportCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = self._tmp.name

    def tearDown(self):
        self._tmp.cleanup()

    def export(self, files, name="course.imscc", sub=""):
        folder = os.path.join(self.dir, sub)
        os.makedirs(folder, exist_ok=True)
        path = os.path.join(folder, name)
        with zipfile.ZipFile(path, "w") as z:
            for fn, text in files.items():
                z.writestr(fn, text)
        return path

    def analyse(self, items, threshold=85.0, full=False):
        """Put the items in one bank and return (variants, clusters), or
        (variants, clusters, lookalikes) with full=True."""
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", items)})
        questions, _, _, _ = qd.collect_questions([path])
        variants = qd.build_variants(questions)
        clusters, _, lookalikes = qd.find_version_clusters(variants, threshold)
        return (variants, clusters, lookalikes) if full else (variants, clusters)


CHOICES = ["The vocal qualities that accompany speech", "The study of personal space",
           "Communication through touch", "The use of time to communicate"]
STEM = "Which of the following best describes paralanguage?"


class Parsing(ExportCase):
    def test_key_is_read_per_question_type(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, STEM, CHOICES, [CHOICES[0]]),
            item("multiple_answers_question", "Pick two.", ["A one", "B two", "C three"], ["A one", "B two"]),
            item("short_answer_question", "Personal space is ____.", answers=["proxemics"]),
            item("matching_question", "Match.", rows=[("Touch", "Haptics"), ("Space", "Proxemics")]),
            item("essay_question", "Discuss."),
        ])})
        qs, _, _, _ = qd.collect_questions([path])
        by_type = {q.question_type: q for q in qs}
        self.assertEqual(by_type[MC].correct, [CHOICES[0]])
        self.assertEqual(sorted(by_type["multiple_answers_question"].correct), ["A one", "B two"])
        self.assertEqual(by_type["short_answer_question"].correct, ["proxemics"])
        self.assertEqual(sorted(by_type["matching_question"].key_pairs), [("Space", "Proxemics"), ("Touch", "Haptics")])
        self.assertEqual(by_type["essay_question"].correct, [])

    def test_image_only_stems_are_kept_and_distinguished(self):
        img = '<img src="$IMS-CC-FILEBASE$/Uploaded%20Media/{}.png?canvas_download=1" alt="x">'
        variants, clusters = self.analyse([
            item(MC, img.format("chart one"), CHOICES, [CHOICES[0]]),
            item(MC, img.format("diagram"), CHOICES, [CHOICES[0]]),
        ])
        self.assertEqual(len(variants), 2)
        self.assertIn("[image: chart one.png]", variants[0].rep.text)

    def test_same_file_name_exports_do_not_mix(self):
        a = self.export({"q/ga.xml": quiz("ga", "Quiz A", [item(MC, STEM, CHOICES, [CHOICES[0]])])}, sub="fall")
        b = self.export({"q/gb.xml": quiz("gb", "Quiz B", [item("essay_question", "Discuss.")])}, sub="spring")
        qs, _, _, _ = qd.collect_questions([a, b])
        self.assertEqual(len(qs), 2)
        self.assertEqual(len({q.source_file for q in qs}), 2)

    def test_bank_usage_and_unresolved_references(self):
        path = self.export({
            "q/gq.xml": quiz("gq", "Midterm", [], bank_refs=["gb", "gmissing"]),
            "non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [item("essay_question", "Discuss.")]),
        })
        _, unresolved, usage, _ = qd.collect_questions([path])
        self.assertEqual([u[2] for u in unresolved], ["gmissing"])
        self.assertEqual(list(usage.values()), [{"Midterm"}])

    def test_container_ident_regex(self):
        for tag in ('<assessment ident="g1" title="Q">', '<assessment title="Q" ident="g1">',
                    "<qti:objectbank ident='g1'>"):
            self.assertEqual(qd.CONTAINER_IDENT_RE.search(tag).group(1), "g1")


class Versions(ExportCase):
    def test_reworded_stem_and_reworded_choice_are_clustered(self):
        tweaked = [CHOICES[0].replace("vocal qualities", "vocal qualities, such as pitch,")] + CHOICES[1:]
        _, clusters = self.analyse([
            item(MC, STEM, CHOICES, [CHOICES[0]]),
            item(MC, STEM.replace("best describes", "most accurately describes"), CHOICES, [CHOICES[0]]),
            item(MC, STEM, tweaked, [tweaked[0]]),
        ])
        self.assertEqual([len(c) for c in clusters], [3])
        ref = clusters[0][0]
        self.assertFalse(any(qd.key_differs(ref, v) for v in clusters[0][1:]))

    def test_generic_stem_with_unrelated_choices_is_not_clustered(self):
        _, clusters = self.analyse([
            item(MC, "Which of the following is true?", CHOICES, [CHOICES[0]]),
            item(MC, "Which of the following is true?",
                 ["Ethos means credibility", "Pathos means logic", "Logos means emotion", "Kairos means volume"],
                 ["Ethos means credibility"]),
        ])
        self.assertEqual(clusters, [])

    def test_same_frame_different_content_is_not_clustered(self):
        frame = '"To persuade my audience that {}" is a specific purpose statement for a persuasive speech on a question of'
        kinds = ["fact.", "value.", "policy.", "opinion.", "attitude."]
        _, clusters, lookalikes = self.analyse([
            item(MC, frame.format("long-term exposure to electromagnetic fields can cause serious health problems"), kinds, ["fact."]),
            item(MC, frame.format("soccer will become the highest revenue-producing sport in the United States"), kinds, ["fact."]),
            item(MC, frame.format("eating meat is unethical"), kinds, ["value."]),
        ], full=True)
        self.assertEqual((clusters, lookalikes), ([], []))

    def test_appended_hint_is_still_the_same_question(self):
        stem = "According to your textbook, the credibility of a speaker at the end of the speech is called"
        kinds = ["initial credibility", "derived credibility", "terminal credibility", "final credibility"]
        _, clusters = self.analyse([
            item(MC, stem, kinds, [kinds[2]]),
            item(MC, stem + " Warning: only two of these are actual types of credibility.", kinds, [kinds[2]]),
        ])
        self.assertEqual([len(c) for c in clusters], [1 + 1])

    def test_exact_copies_collapse_to_one_variant(self):
        variants, clusters = self.analyse([item(MC, STEM, CHOICES, [CHOICES[0]]) for _ in range(3)])
        self.assertEqual((len(variants), len(variants[0].occurrences), clusters), (1, 3, []))

    def test_question_type_change_is_a_separate_version(self):
        _, clusters = self.analyse([
            item(MC, STEM, CHOICES, [CHOICES[0]]),
            item("multiple_answers_question", STEM, CHOICES, [CHOICES[0]]),
        ])
        self.assertEqual([len(c) for c in clusters], [2])

    def test_threshold_tightens_every_rule(self):
        items = [
            item(MC, "Which term names the vocal qualities that accompany a spoken message?", CHOICES, [CHOICES[0]]),
            item(MC, "Which word names the vocal qualities accompanying spoken messages?",
                 [CHOICES[0], "The study of space", "Communicating by touch", CHOICES[3]], [CHOICES[0]]),
        ]
        self.assertEqual(len(self.analyse(items, 85)[1]), 1)
        # stem and choices are both in the high 80s: grouped at the default,
        # not once the threshold is raised past them
        self.assertEqual(len(self.analyse(items, 99)[1]), 0)


class AnswerKey(ExportCase):
    def key_changed(self, a, b):
        _, clusters = self.analyse([a, b])
        self.assertEqual(len(clusters), 1, "expected the two questions to be clustered")
        return qd.key_differs(*clusters[0])

    def test_key_moved_to_another_choice(self):
        self.assertTrue(self.key_changed(item(MC, STEM, CHOICES, [CHOICES[0]]), item(MC, STEM, CHOICES, [CHOICES[1]])))

    def test_reworded_credited_choice_is_not_a_key_change(self):
        new = ["The vocal qualities which accompany speech"] + CHOICES[1:]
        self.assertFalse(self.key_changed(item(MC, STEM, CHOICES, [CHOICES[0]]), item(MC, STEM, new, [new[0]])))

    def test_newly_added_credited_choice_is_a_key_change(self):
        new = CHOICES + ["The vocal qualities that replace speech"]
        self.assertTrue(self.key_changed(item(MC, STEM, CHOICES, [CHOICES[0]]), item(MC, STEM, new, [new[-1]])))

    def test_negated_true_false_is_a_lookalike_not_a_version(self):
        _, clusters, lookalikes = self.analyse([
            item("true_false_question", "Kinesics is the study of body movement.", ["True", "False"], ["True"]),
            item("true_false_question", "Kinesics is not the study of body movement.", ["True", "False"], ["False"]),
        ], full=True)
        self.assertEqual(clusters, [])
        self.assertEqual([len(g) for g in lookalikes], [2])

    def test_reworded_stem_with_changed_answer_is_a_lookalike(self):
        _, clusters, lookalikes = self.analyse([
            item(MC, "Which of the following would you most likely find in a speech conclusion?", CHOICES, [CHOICES[0]]),
            item(MC, "Which of the following would you least likely find in a speech conclusion?", CHOICES, [CHOICES[1]]),
        ], full=True)
        self.assertEqual(clusters, [])
        self.assertEqual([len(g) for g in lookalikes], [2])

    def test_matching_reworded_prompt_is_not_a_key_change(self):
        rows = [("Communication through touch", "Haptics"), ("Use of personal space", "Proxemics"),
                ("Use of time", "Chronemics"), ("Body movement", "Kinesics")]
        reworded = [("Communicating through touch", "Haptics")] + rows[1:]
        self.assertFalse(self.key_changed(item("matching_question", "Match each term.", rows=rows),
                                          item("matching_question", "Match each term.", rows=reworded)))

    def test_matching_swapped_answers_is_a_key_change(self):
        rows = [("Communication through touch", "Haptics"), ("Use of personal space", "Proxemics"),
                ("Use of time", "Chronemics"), ("Body movement", "Kinesics")]
        swapped = [(rows[0][0], "Proxemics"), (rows[1][0], "Haptics")] + rows[2:]
        self.assertTrue(self.key_changed(item("matching_question", "Match each term.", rows=rows),
                                         item("matching_question", "Match each term.", rows=swapped)))


class ReviewLists(ExportCase):
    def test_pair_joined_through_a_third_version_is_not_listed_twice(self):
        # A and C: reworded, same answer. C and B: same wording, different key.
        # So A, B, C form one cluster, and A/B (reworded + different answer)
        # must not also appear as a look-alike group.
        reworded = STEM.replace("best describes", "most accurately describes")
        _, clusters, lookalikes = self.analyse([
            item(MC, STEM, CHOICES, [CHOICES[0]]),
            item(MC, reworded, CHOICES, [CHOICES[1]]),
            item(MC, reworded, CHOICES, [CHOICES[0]]),
        ], full=True)
        self.assertEqual([len(c) for c in clusters], [3])
        self.assertEqual(lookalikes, [])

    def test_lookalike_groups_do_not_chain(self):
        usage = [5, 1, 1, 1, 1]
        self.assertEqual(qd.group_lookalikes([(0, 1), (1, 2), (2, 3)], usage), [[0, 1], [2, 3]])
        # 4's only look-alike (1) is taken, so it joins 1's group and is not dropped
        self.assertEqual(qd.group_lookalikes([(0, 1), (1, 4)], usage), [[0, 1, 4]])
        self.assertEqual(qd.group_lookalikes([], usage), [])

    def test_unrelated_choices_under_a_similar_stem_are_not_linked(self):
        frame = "Which of the following is presented in your textbook as a guideline for {}?"
        ethics = ["Explain your motives for speaking.", "Cite your sources in the introduction.",
                  "Make sure your goals are ethically sound.", "All of these answers are correct.",
                  "Both explain your motives for speaking and cite your sources."]
        visual = ["Display visual aids where listeners can see them.", "Pass small visual aids among the audience.",
                  "Display visual aids only while discussing them.", "All of these answers are correct.",
                  "Both display visual aids where listeners can see them and only while discussing them."]
        _, clusters, lookalikes = self.analyse([
            item(MC, frame.format("ethical speechmaking"), ethics, [ethics[2]]),
            item(MC, frame.format("presenting visual aids"), visual, [visual[4]]),
        ], full=True)
        self.assertEqual((clusters, lookalikes), ([], []))

    def test_same_stem_list_shows_each_cluster_once(self):
        other = ["Ethos means credibility", "Pathos means logic", "Logos means emotion", "Kairos means volume"]
        tweaked = CHOICES[:3] + ["The use of time to send messages"]
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, "Which of the following is true?", CHOICES, [CHOICES[0]]),
            item(MC, "Which of the following is true?", tweaked, [CHOICES[0]]),
            item(MC, "Which of the following is true?", other, [other[0]]),
        ])})
        variants = qd.build_variants(qd.collect_questions([path])[0])
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        self.assertEqual([len(g) for g in qd.find_same_stem_unrelated(variants, roots)], [2])

    def test_same_choices_list_needs_the_same_answer(self):
        kinds = ["pseudolistening", "selective listening", "defensive listening", "ambushing"]
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, "Giving only the appearance of being attentive is termed", kinds, [kinds[0]]),
            item(MC, "What is it called when you look like you are paying attention but are not?", kinds, [kinds[0]]),
            item(MC, "Focusing only on particular parts of a message is called", kinds, [kinds[1]]),
        ])})
        variants = qd.build_variants(qd.collect_questions([path])[0])
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        groups = qd.find_same_choices_other_stem(variants, roots)
        self.assertEqual([len(g) for g in groups], [2])
        self.assertTrue(all(v.correct == frozenset({"pseudolistening"}) for v in groups[0]))

    def test_repeated_choice_found_even_when_a_clean_copy_comes_first(self):
        variants, _ = self.analyse([
            item(MC, STEM, CHOICES, [CHOICES[0]]),
            item(MC, STEM, CHOICES + [CHOICES[1] + " "], [CHOICES[0]]),
        ])
        repeated = qd.find_repeated_choices(variants)
        self.assertEqual([len(v.occurrences) for v in repeated], [1])
        self.assertEqual(len(repeated[0].rep.choices), 5)


class WholeContainers(ExportCase):
    # stems with nothing in common, so no two are mistaken for versions
    STEMS = {
        "ethos": "Explain how a speaker builds credibility with a skeptical audience.",
        "pathos": "Describe one risk of relying on emotional appeals in a eulogy.",
        "logos": "What makes an argument from statistics convincing to listeners?",
        "kinesics": "Give two examples of gestures that changed the meaning of a spoken sentence.",
        "proxemics": "Why does standing too close to a stranger feel uncomfortable?",
        "haptics": "When is touch appropriate in a professional greeting?",
        "delivery": "Compare reading from a manuscript with speaking from brief notes.",
        "outlining": "List the symbols used for main points and subpoints in a formal outline.",
    }
    TOPICS = ["ethos", "pathos", "logos", "kinesics", "proxemics", "haptics"]

    def q(self, topic, reworded=False):
        stem = self.STEMS[topic]
        return item("essay_question", stem.replace("greeting?", "greeting today?") if reworded else stem)

    def compare(self, files, overlap=80.0):
        path = self.export(files)
        variants = qd.build_variants(qd.collect_questions([path])[0])
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        identical, overlaps = qd.find_duplicate_containers(variants, roots, overlap)
        return ([sorted(c.title for c in g) for g in identical],
                {c.title: [(o.title, exact, modified) for o, exact, modified, _ in hits] for c, hits in overlaps})

    def test_identical_bank_and_quiz(self):
        five = self.TOPICS[:5]
        identical, overlaps = self.compare({
            "non_cc_assessments/g1.xml.qti": bank("g1", "Bank A", [self.q(t) for t in five]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "Bank B", [self.q(t) for t in reversed(five)] + [self.q(five[0])]),
            "q/g3.xml": quiz("g3", "Quiz C", [self.q(t) for t in self.TOPICS[3:]]),
        })
        self.assertEqual(identical, [["Bank A", "Bank B"]])  # order and an internal repeat don't matter
        self.assertEqual(overlaps, {})

    def test_subset_is_fully_contained(self):
        identical, overlaps = self.compare({
            "non_cc_assessments/g1.xml.qti": bank("g1", "Small", [self.q(t) for t in self.TOPICS[:4]]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "Large", [self.q(t) for t in self.TOPICS]),
        })
        self.assertEqual(identical, [])
        self.assertEqual(overlaps, {"Small": [("Large", 4, 0)]})

    def test_modified_version_counts_as_a_match(self):
        _, overlaps = self.compare({
            "non_cc_assessments/g1.xml.qti": bank("g1", "Small", [self.q(t) for t in self.TOPICS[:3]] + [self.q("haptics", reworded=True)]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "Large", [self.q(t) for t in self.TOPICS]),
        })
        self.assertEqual(overlaps, {"Small": [("Large", 3, 1)]})

    def test_partial_overlap_respects_the_setting(self):
        files = {
            "non_cc_assessments/g1.xml.qti": bank("g1", "Half", [self.q(t) for t in self.TOPICS[:2]] + [self.q("delivery"), self.q("outlining")]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "Large", [self.q(t) for t in self.TOPICS]),
        }
        self.assertEqual(self.compare(files)[1], {})
        self.assertEqual(self.compare(files, overlap=50.0)[1], {"Half": [("Large", 2, 0)]})

    def test_tiny_containers_are_skipped(self):
        identical, overlaps = self.compare({
            "non_cc_assessments/g1.xml.qti": bank("g1", "Tiny A", [self.q("ethos"), self.q("pathos")]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "Tiny B", [self.q("ethos"), self.q("pathos")]),
        })
        self.assertEqual((identical, overlaps), ([], {}))


class NewQuizzes(ExportCase):
    """New Quizzes as a course export writes them: the same QTI as Classic
    Quizzes, with a marker on the quiz, on the item bank and on each question."""
    PROMPTS = list(WholeContainers.STEMS.values())

    def essays(self, picks, new=True):
        return [item("essay_question", self.PROMPTS[i], new=new) for i in picks]

    def test_new_quizzes_and_item_banks_are_recognised(self):
        path = self.export({
            "non_cc_assessments/g1.xml.qti": quiz("g1", "A", self.essays([0]), new=True),
            "non_cc_assessments/g2.xml.qti": bank("g2", "B", self.essays([1]), new=True),
            "non_cc_assessments/g3.xml.qti": quiz("g3", "C", self.essays([2], new=False)),
            "non_cc_assessments/g4.xml.qti": bank("g4", "D", self.essays([3], new=False)),
        })
        qs, _, _, containers = qd.collect_questions([path])
        self.assertEqual({c.title: qd.kind_label(c.kind, c.engine) for c in containers},
                         {"A": "new quiz", "B": "item bank", "C": "quiz", "D": "bank"})
        self.assertEqual({q.quiz_title: q.engine for q in qs},
                         {"A": "new", "B": "new", "C": "classic", "D": "classic"})

    def test_unmarked_quiz_is_recognised_by_its_questions(self):
        path = self.export({"q/g1.xml": quiz("g1", "Quiz", self.essays([0]))})
        self.assertEqual([c.engine for c in qd.collect_questions([path])[3]], ["new"])

    def test_same_question_in_classic_and_new_is_one_question(self):
        path = self.export({
            "non_cc_assessments/g1.xml.qti": bank("g1", "Classic bank", [item(MC, STEM, CHOICES, [CHOICES[0]])]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "Item bank", [item(MC, STEM, CHOICES, [CHOICES[0]], new=True)],
                                                  new=True),
        })
        variants = qd.build_variants(qd.collect_questions([path])[0])
        self.assertEqual([sorted(q.engine for q in v.occurrences) for v in variants], [["classic", "new"]])
        self.assertEqual([q.correct for q in variants[0].occurrences], [[CHOICES[0]], [CHOICES[0]]])

    def test_quiz_that_only_pulls_from_banks_is_kept(self):
        path = self.export({
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Exam", [], draws=[("gb", 5), ("gmissing", 2)], new=True),
            "non_cc_assessments/gb.xml.qti": bank("gb", "Chapter 1", self.essays([0, 1, 2]), new=True),
        })
        qs, unresolved, usage, containers = qd.collect_questions([path])
        exam = next(c for c in containers if c.title == "Exam")
        self.assertEqual((exam.kind, exam.engine, exam.draws), ("quiz", "new", [("gb", 5), ("gmissing", 2)]))
        self.assertEqual({q.quiz_title for q in qs}, {"Chapter 1"})
        self.assertEqual([u[2] for u in unresolved], ["gmissing"])
        self.assertEqual(list(usage.values()), [{"Exam"}])
        self.assertEqual(qd.describe_contents(exam), "draws 7 from 2 banks")

    def test_quizzes_pulling_the_same_from_the_same_banks_are_identical(self):
        path = self.export({
            "non_cc_assessments/ga.xml.qti": bank("ga", "Chapter 1", self.essays([0, 1, 2]), new=True),
            "non_cc_assessments/gb.xml.qti": bank("gb", "Chapter 1 copy", self.essays([2, 1, 0]), new=True),
            "non_cc_assessments/gc.xml.qti": bank("gc", "Chapter 2", self.essays([3, 4, 5]), new=True),
            "non_cc_assessments/q1.xml.qti": quiz("q1", "Exam", [], draws=[("ga", 2)], new=True),
            "non_cc_assessments/q2.xml.qti": quiz("q2", "Exam copy", [], draws=[("gb", 2)], new=True),
            "non_cc_assessments/q3.xml.qti": quiz("q3", "Longer exam", [], draws=[("ga", 3)], new=True),
            "non_cc_assessments/q4.xml.qti": quiz("q4", "Other exam", [], draws=[("gc", 2)], new=True),
        })
        qs, _, _, containers = qd.collect_questions([path])
        variants = qd.build_variants(qs)
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        identical, overlaps = qd.find_duplicate_containers(variants, roots, 80.0, containers)
        self.assertEqual([sorted(c.title for c in g) for g in identical],
                         [["Chapter 1", "Chapter 1 copy"], ["Exam", "Exam copy"]])
        self.assertEqual(overlaps, [])

    def test_single_questions_used_from_a_bank(self):
        qs = self.essays([0, 1, 2])
        a, b = item_id(qs[0]), item_id(qs[1])
        path = self.export({
            "non_cc_assessments/gb.xml.qti": bank("gb", "Midterm pool", qs, new=True),
            "non_cc_assessments/q1.xml.qti": quiz("q1", "Exam", [], picks=[("gb", a), ("gb", b)], new=True),
            "non_cc_assessments/q2.xml.qti": quiz("q2", "Exam copy", [], picks=[("gb", b), ("gb", a)], new=True),
            "non_cc_assessments/q3.xml.qti": quiz("q3", "Other exam", [], picks=[("gb", a), ("gmissing", "x")], new=True),
            "non_cc_assessments/q4.xml.qti": quiz("q4", "All of it", [], draws=[("gb", None)], new=True),
            "non_cc_assessments/q5.xml.qti": quiz("q5", "Nothing", [], new=True),
        })
        questions, unresolved, _, containers = qd.collect_questions([path])
        self.assertEqual(len(questions), 3)  # the picks are references, not extra copies
        self.assertEqual([u[1:3] for u in unresolved], [("Other exam", "gmissing")])
        variants = qd.build_variants(questions)
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        identical, overlaps = qd.find_duplicate_containers(variants, roots, 80.0, containers)
        self.assertEqual([sorted(c.title for c in g) for g in identical], [["Exam", "Exam copy"]])
        self.assertEqual(overlaps, [])
        by_title = {c.title: c for c in containers}
        self.assertEqual(qd.describe_contents(by_title["Exam"]), "uses 2 bank questions")
        self.assertEqual(len(by_title["Other exam"].picked), 1)
        self.assertEqual(by_title["All of it"].draws, [("gb", None)])
        self.assertEqual(qd.describe_contents(by_title["All of it"]), "uses all of 1 bank")

        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        self.assertIn("2 chosen questions from [item bank] Midterm pool", report)
        self.assertIn("1 chosen question could not be found in the export", report)
        self.assertIn("all questions from [item bank] Midterm pool", report)
        self.assertIn("Quizzes with nothing in the export", report)
        self.assertIn("[new quiz] <strong>Nothing</strong>", report)

    def test_missing_picks_are_counted_by_pick_not_by_question(self):
        # two picks of identically worded questions, plus the same missing pick twice
        twins = [item("essay_question", self.PROMPTS[0], new=True), item("essay_question", self.PROMPTS[0], new=True)]
        path = self.export({
            "non_cc_assessments/gb.xml.qti": bank("gb", "Pool", twins + self.essays([1, 2]), new=True),
            "non_cc_assessments/q1.xml.qti": quiz("q1", "Exam", [], new=True, picks=[
                ("gb", item_id(twins[0])), ("gb", item_id(twins[1])), ("gb", "gone"), ("gb", "gone")]),
        })
        questions, _, _, containers = qd.collect_questions([path])
        variants = qd.build_variants(questions)
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        qd.find_duplicate_containers(variants, roots, 80.0, containers)
        exam = next(c for c in containers if c.title == "Exam")
        self.assertEqual((len(exam.picked), exam.lost_picks), (1, 2))

    def test_very_low_threshold_does_not_crash(self):
        _, clusters = self.analyse([item(MC, STEM, CHOICES, [CHOICES[0]]),
                                    item(MC, STEM + " Explain.", CHOICES, [CHOICES[0]])], threshold=5.0)
        self.assertEqual([len(c) for c in clusters], [2])

    def test_report_names_new_quizzes_and_lists_what_they_pull(self):
        path = self.export({
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Exam", [], draws=[("gb", 2)], new=True),
            "non_cc_assessments/gb.xml.qti": bank("gb", "Chapter 1", self.essays([0, 1, 2]), new=True),
            "non_cc_assessments/gc.xml.qti": bank("gc", "Old bank", self.essays([0, 1, 2], new=False)),
        })
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        self.assertIn("1 quiz and 2 banks (from New Quizzes: 1 quiz and 1 item bank)", report)
        self.assertIn("[new quiz] <strong>Exam</strong>", report)
        self.assertIn("2 at random from [item bank] Chapter 1", report)
        # the classic bank and the item bank hold the same three questions
        self.assertIn("2 with the same 3 questions", report)
        self.assertIn("[item bank] <strong>Chapter 1</strong>", report)
        self.assertIn("[bank] <strong>Old bank</strong>", report)


class Outputs(ExportCase):
    def run_main(self, *extra):
        tweaked = CHOICES[:3] + ["The use of time to send messages"]
        path = self.export({
            "non_cc_assessments/gb.xml.qti": bank("gb", "Pool", [
                item(MC, STEM, CHOICES, [CHOICES[0]], new=True),
                item(MC, STEM, tweaked, [CHOICES[1]], new=True),
                item("essay_question", "Discuss ethos.", new=True)], new=True),
            "non_cc_assessments/gc.xml.qti": bank("gc", "Old pool", [item(MC, STEM, CHOICES, [CHOICES[0]])]),
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Exam", [], draws=[("gb", None), ("gmissing", 2)], new=True),
        })
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, *extra]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            return f.read()

    def test_json_has_the_report_sections(self):
        target = os.path.join(self.dir, "report.json")
        self.run_main("--json", target, "-t", "80")
        with open(target, encoding="utf-8") as f:
            data = json.load(f)
        self.assertEqual(data["settings"]["threshold"], 80.0)
        self.assertEqual([s["export"] for s in data["settings"]["sources"]], ["course.imscc"])
        self.assertEqual({a["export"] for q in data["exact_duplicates"] for a in q["locations"]}, {"course.imscc"})
        self.assertEqual(data["summary"], {"questions": 4, "distinct_questions": 3, "quizzes": 1, "banks": 2,
                                           "new_quizzes": 1, "item_banks": 1})
        self.assertEqual(data["question_images"], {"compared_by_content": 0, "compared_by_file_name": 0,
                                                   "compared_by_web_address": 0})
        exam = next(c for c in data["quizzes_and_banks"] if c["title"] == "Exam")
        self.assertEqual(exam["kind"], "new quiz")
        self.assertEqual(exam["draws"], [{"bank_id": "gb", "bank_title": "Pool", "count": None},
                                         {"bank_id": "gmissing", "bank_title": None, "count": 2}])
        self.assertEqual([u["bank_id"] for u in data["unresolved_bank_references"]], ["gmissing"])
        (cluster,) = data["modified_questions"]
        self.assertTrue(cluster["answer_key_differs"])
        a, b = cluster["versions"]
        self.assertEqual((a["version"], a["changes_from_a"], a["answer_key"]), ("A", None, [CHOICES[0]]))
        self.assertEqual(sorted(p["kind"] for p in a["locations"]), ["bank", "item bank"])
        self.assertEqual(a["locations"][0]["drawn_on_by"], ["Exam"])
        self.assertTrue(b["changes_from_a"]["choices_changed"] and b["changes_from_a"]["answer_key_differs"])
        self.assertFalse(b["changes_from_a"]["stem_reworded"])
        self.assertEqual([q["stem"] for q in data["exact_duplicates"]], [STEM])

    def test_no_json_unless_asked(self):
        self.run_main()
        self.assertEqual(sorted(os.listdir(self.dir)), ["course.imscc", "report.html"])

    def test_contents_links_point_at_real_sections(self):
        import re
        report = self.run_main()
        nav = report[report.index('<nav class="toc"'):report.index("</nav>")]
        links = re.findall(r"href='#([^']+)'", nav)
        self.assertNotIn("<span class='src'></span>", report)
        self.assertIn("modified", links)
        self.assertIn("unresolved", links)
        for anchor in links:
            self.assertEqual(report.count(f'id="{anchor}"'), 1, anchor)


class OrphanBanks(ExportCase):
    """The bank usage and orphan check: which banks quizzes use, and for the
    rest, whether every question also lives somewhere in use."""
    Q1 = (MC, "What is the specific purpose of a speech?",
          ["The goal in one sentence", "The topic area", "The audience", "The thesis outline"],
          ["The goal in one sentence"])
    Q2 = (MC, STEM, CHOICES, [CHOICES[0]])
    Q2_TWEAKED = (MC, STEM, CHOICES[:3] + ["The use of time to send messages"], [CHOICES[0]])
    Q3 = (MC, "Which appeal relies on the credibility of the speaker?",
          ["Ethos", "Pathos", "Logos", "Kairos"], ["Ethos"])
    Q4 = (MC, "Monroe's motivated sequence begins with which step?",
          ["Attention", "Need", "Satisfaction", "Action"], ["Attention"])
    Q5 = (MC, "A fallacy that attacks the person instead of the argument is called what?",
          ["Ad hominem", "Straw man", "Slippery slope", "Bandwagon"], ["Ad hominem"])

    def course(self):
        q5 = item(*self.Q5)
        return {
            "non_cc_assessments/ga.xml.qti": bank("ga", "Persuasion", [item(*self.Q1), item(*self.Q2)]),
            "non_cc_assessments/gb.xml.qti": bank("gb", "Persuasion", [item(*self.Q1), item(*self.Q2)]),
            "non_cc_assessments/gc.xml.qti": bank("gc", "Persuasion", [item(*self.Q1), item(*self.Q2_TWEAKED)]),
            "non_cc_assessments/gd.xml.qti": bank("gd", "Old final", [item(*self.Q3)]),
            "non_cc_assessments/ge.xml.qti": bank("ge", "Old final", [item(*self.Q3)]),
            "non_cc_assessments/gf.xml.qti": bank("gf", "Extra", [item(*self.Q4), q5]),
            "non_cc_assessments/gg.xml.qti": bank("gg", "Empty", []),
            "non_cc_assessments/gq1.xml.qti": quiz("gq1", "Unit 3 quiz", [], bank_refs=["ga"]),
            "non_cc_assessments/gq2.xml.qti": quiz("gq2", "Fallacies quiz", [item(*self.Q5, bank_ref=item_id(q5))]),
        }

    def check(self, files=None):
        path = self.export(files or self.course())
        questions, _, _, containers = qd.collect_questions([path])
        usage, seen, resolved = qd.find_bank_usage(questions, containers)
        variants = qd.build_variants(questions)
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        qd.find_duplicate_containers(variants, roots, containers=containers)
        return {b.bank.cid: b for b in qd.check_orphan_banks(variants, roots, containers, usage)}, seen, resolved

    def test_usage_status(self):
        audit, seen, resolved = self.check()
        self.assertEqual((seen, resolved), (1, 1))
        self.assertEqual({k: b.status for k, b in audit.items()},
                         {"ga": "linked", "gb": "unused", "gc": "unused", "gd": "unused", "ge": "unused",
                          "gf": "copied only", "gg": "unused"})
        self.assertEqual(audit["ga"].usage.linked_by, {"Unit 3 quiz"})
        self.assertEqual(audit["gf"].usage.copied_into, {"Fallacies quiz"})
        self.assertEqual(audit["ga"].verdict, "")  # linked banks get no verdict

    def test_verdicts(self):
        audit, _, _ = self.check()
        self.assertEqual({k: b.verdict for k, b in audit.items() if b.verdict},
                         {"gb": "safe", "gc": "review", "gd": "keep_copy", "ge": "covered", "gf": "unique",
                          "gg": "empty"})
        self.assertEqual(audit["gc"].count("modified"), 1)
        modified = next(f for f in audit["gc"].findings if f.coverage == "modified")
        self.assertEqual(modified.places, ["Persuasion"])
        # Q4 is unique; Q5 was copied into a quiz, so it is in use as is
        self.assertEqual((audit["gf"].count("unique"), audit["gf"].count("exact")), (1, 1))

    def test_unlinked_copies_do_not_vouch_for_each_other_but_one_is_kept(self):
        audit, _, _ = self.check()
        (kept,) = audit["gd"].findings
        (other,) = audit["ge"].findings
        self.assertEqual((kept.coverage, kept.places), ("unused_only", ["Old final"]))
        self.assertEqual((other.coverage, other.places), ("kept_copy", ["Old final"]))
        self.assertEqual(audit["gd"].kept_for, [audit["ge"]])
        self.assertEqual(audit["ge"].relies_on, [audit["gd"]])

    def test_keep_the_bank_holding_the_most(self):
        q = [item(*self.Q3), item(*self.Q4), item(*self.Q5)]
        files = {
            "non_cc_assessments/g1.xml.qti": bank("g1", "Small", [item(*self.Q3)]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "Big", q),
            "non_cc_assessments/g3.xml.qti": bank("g3", "Medium", [item(*self.Q4), item(*self.Q5)]),
        }
        audit, _, _ = self.check(files)
        self.assertEqual({k: b.verdict for k, b in audit.items()},
                         {"g1": "covered", "g2": "keep_copy", "g3": "covered"})

    def test_bank_with_unique_questions_is_kept_first(self):
        # g2 must stay for its unique Q1, and it also holds Q3, so g1 (only Q3) can go
        files = {
            "non_cc_assessments/g1.xml.qti": bank("g1", "A copy", [item(*self.Q3)]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "B mixed", [item(*self.Q1), item(*self.Q3)]),
        }
        audit, _, _ = self.check(files)
        self.assertEqual((audit["g1"].verdict, audit["g2"].verdict), ("covered", "unique"))

    def test_kept_bank_with_another_version_means_review(self):
        files = {
            "non_cc_assessments/g1.xml.qti": bank("g1", "A", [item(*self.Q2), item(*self.Q3)]),
            "non_cc_assessments/g2.xml.qti": bank("g2", "B", [item(*self.Q2_TWEAKED)]),
        }
        audit, _, _ = self.check(files)
        self.assertEqual((audit["g1"].verdict, audit["g2"].verdict), ("unique", "review"))
        (finding,) = audit["g2"].findings
        self.assertEqual(finding.coverage, "kept_version")

    def test_unresolved_back_pointer_is_counted_but_not_used(self):
        files = self.course()
        files["non_cc_assessments/gq3.xml.qti"] = quiz("gq3", "Inline quiz", [item(*self.Q4, bank_ref="gnothere")])
        audit, seen, resolved = self.check(files)
        self.assertEqual((seen, resolved), (2, 1))
        # Q4 now has an in-use copy in a quiz, so the Extra bank is safe to delete
        self.assertEqual((audit["gf"].status, audit["gf"].verdict), ("copied only", "safe"))

    def run_main(self, *extra):
        path = self.export(self.course())
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "report.json")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--json", target, *extra]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f, open(target, encoding="utf-8") as g:
            return f.read(), json.load(g)

    def test_excluded_quiz_still_counts_as_using_its_bank(self):
        _, data = self.run_main("--exclude-title", "Unit 3")
        rows = {b["id"]: b for b in data["bank_usage_audit"]}
        self.assertEqual(rows["ga"]["status"], "linked")
        self.assertEqual(rows["ga"]["linked_by"], ["Unit 3 quiz"])

    def test_excluded_type_blocks_safe(self):
        _, data = self.run_main("--exclude-type", MC)
        rows = {b["id"]: b for b in data["bank_usage_audit"]}
        self.assertEqual((rows["gb"]["verdict"], rows["gb"]["counts"]["not_compared"]), ("review", 2))

    def test_summary_names_banks_and_folds_numbered_families(self):
        files = {f"non_cc_assessments/gc{k}.xml.qti": bank(f"gc{k}", f"chapter_{k}-_topic", [item(*q)])
                 for k, q in enumerate((self.Q1, self.Q3, self.Q4), 1)}
        files["non_cc_assessments/gx.xml.qti"] = bank("gx", "Lone bank", [item(*self.Q5)])
        path = self.export(files)
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        summary = report[report.index("class='audit-summary'"):]
        summary = summary[:summary.index("</ul>\n")]
        self.assertIn("Has unique questions", summary)
        self.assertIn(qd.VERDICT_HELP[qd.VERDICT_UNIQUE], summary)
        self.assertIn("the 3 \u201cchapter_N-_\u2026\u201d banks: 1 each, all unique", summary)
        self.assertIn("\u201cLone bank\u201d</a>: 1 (all unique)", summary)

    def test_review_badge_is_readable(self):
        report, _ = self.run_main()
        self.assertIn(".verdict.review {{ color: var(--review-fg)".replace("{{", "{"), report)
        self.assertIn("--review-fg: #8a6400", report)

    def test_every_kept_bank_is_counted_in_the_summary(self):
        # nine pairs of unused banks, each pair sharing a question no quiz uses
        questions = [
            ("Ethos rests on what?", ["Credibility", "Emotion", "Logic"]),
            ("Name the first step of Monroe's sequence.", ["Attention", "Need", "Action"]),
            ("Proxemics studies which cue?", ["Space", "Touch", "Time"]),
            ("A thesis statement belongs where?", ["Introduction", "Body", "Appendix"]),
            ("What lowers speech anxiety most reliably?", ["Rehearsal", "Caffeine", "Memorizing jokes"]),
            ("Ad hominem attacks whom?", ["The person", "The data", "The venue"]),
            ("Paraphrasing helps listeners do what?", ["Confirm meaning", "Interrupt", "Daydream"]),
            ("Chronemics concerns which resource?", ["Time", "Color", "Smell"]),
            ("Which visual aid suits trends over years?", ["Line graph", "Pie chart", "Photograph"]),
        ]
        files = {}
        for k, (stem, choices) in enumerate(questions):
            for side in "ab":
                files[f"non_cc_assessments/g{k}{side}.xml.qti"] = bank(
                    f"g{k}{side}", f"Bank {k} {side}", [item(MC, stem, choices, [choices[0]])])
        path = self.export(files)
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        summary = report[report.index("class='audit-summary'"):]
        self.assertIn("<strong>9</strong> <span class='verdict keep_copy'>", summary)
        self.assertIn("<li>and 1 more</li>", summary)

    def test_families_need_a_word_first_and_the_same_punctuation(self):
        qs = [self.Q1, self.Q3, self.Q4, self.Q5, self.Q2,
              (MC, "Which outline lists only key words?", ["Speaking outline", "Formal outline", "Rough outline"],
               ["Speaking outline"])]
        titles = ["101 final", "2018 summer final", "2019 midterm",              # no word first: not a family
                  "Chapter_01_First_Look", "Chapter_04_Words", "chapter_4-_language"]  # two styles: no family of 3
        files = {f"non_cc_assessments/gt{k}.xml.qti": bank(f"gt{k}", t, [item(*q)])
                 for k, (t, q) in enumerate(zip(titles, qs))}
        path = self.export(files)
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        summary = report[report.index("class='audit-summary'"):]
        summary = summary[:summary.index("</ul>\n")]
        self.assertNotIn("\u2026\u201d banks", summary)
        for t in titles:
            self.assertIn(f"\u201c{t}\u201d</a>: 1", summary)

    def test_report_links_review_items_to_their_modified_group(self):
        report, data = self.run_main()
        rows = {b["id"]: b for b in data["bank_usage_audit"]}
        (item_,) = rows["gc"]["questions_to_review"]
        group = item_["modified_group"]
        self.assertEqual(item_["coverage"], "modified")
        self.assertIn(f"href='#modified-{group}'", report)
        self.assertEqual(report.count(f'id="modified-{group}"'), 1)
        self.assertEqual(data["classic_bank_links"], {"quiz_questions_pointing_to_a_bank_question": 1,
                                                      "found_in_export": 1})


def meta(ident, published):
    """A quiz's assessment_meta.xml, the way a course export writes it."""
    return (f'<?xml version="1.0" encoding="UTF-8"?>\n<quiz identifier="{ident}" '
            f'xmlns="http://canvas.instructure.com/xsd/cccv1p0"><title>Quiz</title>'
            f'<available>{"true" if published else "false"}</available></quiz>')


class ReviewFixes2(ExportCase):
    """Fixes from the second outside review."""
    Q1, Q3 = OrphanBanks.Q1, OrphanBanks.Q3

    def main(self, files, *extra):
        path = self.export(files)
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "report.json")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--json", target, *extra]), \
                contextlib.redirect_stdout(io.StringIO()) as printed:
            qd.main()
        with open(out, encoding="utf-8") as f, open(target, encoding="utf-8") as g:
            return f.read(), json.load(g), printed.getvalue()

    def test_minus_after_a_letter_is_kept(self):
        self.assertNotEqual(qd.normalize("x-5"), qd.normalize("x5"))
        self.assertEqual(qd.normalize("x-5"), qd.normalize("x - 5"))
        self.assertEqual(qd.normalize("1-5"), qd.normalize("1 - 5"))

    def test_short_plurals_share_a_word(self):
        self.assertEqual(qd.content_words("cars buses"), qd.content_words("car bus"))

    def test_rewritten_single_row_is_paired_by_elimination(self):
        ref = [("body movement", "kinesics"), ("use of space", "proxemics"), ("touch", "haptics")]
        new = [("body movement", "kinesics"), ("use of space", "proxemics"), ("contact between people", "kinesics")]
        self.assertTrue(qd.pair_key_differs(ref, new))
        same = ref[:2] + [("contact between people", "haptics")]
        self.assertFalse(qd.pair_key_differs(ref, same))

    def test_orphan_check_does_not_need_the_container_comparison(self):
        path = self.export(OrphanBanks().course())
        questions, _, _, containers = qd.collect_questions([path])
        usage, _, _ = qd.find_bank_usage(questions, containers)
        variants = qd.build_variants(questions)
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        audit = {b.bank.cid: b.verdict for b in qd.check_orphan_banks(variants, roots, containers, usage)}
        self.assertEqual(audit["gb"], "safe")
        self.assertEqual(audit["gg"], "empty")

    def test_shared_item_id_credits_every_bank(self):
        q = item(*self.Q1)
        ident = item_id(q)
        files = {"non_cc_assessments/ga.xml.qti": bank("ga", "A", [q]),
                 "non_cc_assessments/gb.xml.qti": bank("gb", "B", [q]),
                 "non_cc_assessments/gq.xml.qti": quiz("gq", "Quiz", [item(*self.Q1, bank_ref=ident)])}
        path = self.export(files)
        questions, _, _, containers = qd.collect_questions([path])
        usage, seen, resolved = qd.find_bank_usage(questions, containers)
        self.assertEqual((seen, resolved), (1, 1))
        self.assertEqual({k[1]: u.status for k, u in usage.items()}, {"ga": "copied only", "gb": "copied only"})

    def test_identical_quizzes_note_excluded_draws(self):
        items = [item(*q) for q in (OrphanBanks.Q1, OrphanBanks.Q3, OrphanBanks.Q4)]
        files = {"non_cc_assessments/gx.xml.qti": bank("gx", "Old pool", [item(*OrphanBanks.Q5)]),
                 "non_cc_assessments/gy.xml.qti": bank("gy", "Other pool", [item(*OrphanBanks.Q2)]),
                 "non_cc_assessments/q1.xml.qti": quiz("q1", "Quiz one", items, draws=[("gx", 1)]),
                 "non_cc_assessments/q2.xml.qti": quiz("q2", "Quiz two", items, draws=[("gy", 1)])}
        report, data, _ = self.main(files, "--exclude-title", "pool")
        self.assertIn("what they take from banks left out of the comparison was not compared", report)
        quiz_one = next(c for c in data["quizzes_and_banks"] if c["title"] == "Quiz one")
        self.assertEqual(quiz_one["draws_from_excluded_banks"], 1)

    def test_published_state_and_only_published(self):
        files = {"non_cc_assessments/ga.xml.qti": bank("ga", "Pool", [item(*self.Q1)]),
                 "non_cc_assessments/gq.xml.qti": quiz("gq", "Old exam", [], bank_refs=["ga"]),
                 "gq/assessment_meta.xml": meta("gq", False)}
        report, data, printed = self.main(files)
        self.assertIn("Published state read for 1 of 1 quiz: 1 unpublished", printed)
        self.assertIn("Old exam <span class='src'>(unpublished)</span>", report)
        (row,) = data["bank_usage_audit"]
        self.assertEqual(row["status"], "linked")
        exam = next(c for c in data["quizzes_and_banks"] if c["title"] == "Old exam")
        self.assertIs(exam["published"], False)
        _, data, _ = self.main(files, "--only-published")
        (row,) = data["bank_usage_audit"]
        self.assertEqual((row["status"], row["verdict"]), ("unused", "unique"))
        self.assertTrue(data["settings"]["only_published"])

    def test_unpublished_quiz_does_not_cover_a_bank_with_only_published(self):
        files = {"non_cc_assessments/ga.xml.qti": bank("ga", "Pool", [item(*self.Q1)]),
                 "non_cc_assessments/gq.xml.qti": quiz("gq", "Old quiz", [item(*self.Q1)]),
                 "gq/assessment_meta.xml": meta("gq", False)}
        _, data, _ = self.main(files)
        self.assertEqual(data["bank_usage_audit"][0]["verdict"], "safe")
        _, data, _ = self.main(files, "--only-published")
        self.assertEqual(data["bank_usage_audit"][0]["verdict"], "unique")

    def test_advanced_settings_reach_the_matching(self):
        items = [item(*q) for q in (OrphanBanks.Q1, OrphanBanks.Q3)]
        files = {"non_cc_assessments/ga.xml.qti": bank("ga", "A", items),
                 "non_cc_assessments/gb.xml.qti": bank("gb", "B", [item(*q) for q in (OrphanBanks.Q1, OrphanBanks.Q3)])}
        _, data, _ = self.main(files)
        self.assertEqual(data["identical_quizzes_and_banks"], [])  # two questions: below the default of 3
        _, data, _ = self.main(files, "--min-container", "2")
        self.assertEqual(len(data["identical_quizzes_and_banks"]), 1)
        self.assertEqual(data["settings"]["min_container"], 2)
        _, data, _ = self.main(files)  # the default comes back on the next run
        self.assertEqual(data["identical_quizzes_and_banks"], [])


class ReviewFixes3(ExportCase):
    """Fixes from the third outside review."""

    def test_math_signs_keep_questions_apart(self):
        pairs = [("x < 3", "x > 3"), ("x <= 3", "x < 3"), ("x = 3", "x 3"), ("1/2", "12"),
                 ("a+b", "ab"), ("50%", "50"), ("2*3", "23"), ("5\u22123", "53"), ("x \u2260 3", "x = 3")]
        for a, b in pairs:
            self.assertNotEqual(qd.normalize(a), qd.normalize(b), (a, b))
        self.assertEqual(qd.normalize("50%"), qd.normalize("50 percent"))
        self.assertEqual(qd.normalize("1/2"), qd.normalize("1 / 2"))
        self.assertEqual(qd.normalize("*Hint:"), qd.normalize("Hint"))

    def test_mixed_draws_are_described(self):
        c = qd.Container("e", "quiz", "q", "Quiz", draws=[("b1", None), ("b2", 5), ("b3", 2)])
        self.assertEqual(qd.describe_contents(c), "uses all of 1 bank + draws 7 from 2 banks")

    def test_items_of_a_nested_bank_belong_to_it_only(self):
        inner = bank("gin", "Inner", [item(*OrphanBanks.Q3)]).split("?>", 1)[1]
        inner = inner.replace(f'<questestinterop xmlns="{NS}">', "").replace("</questestinterop>", "")
        outer = quiz("gq", "Outer", [item(*OrphanBanks.Q1)]).replace("</section></assessment>",
                                                                   f"{inner}</section></assessment>")
        questions, _, _, containers = qd.collect_questions([self.export({"non_cc_assessments/gq.xml.qti": outer})])
        self.assertEqual(sorted((q.quiz_id, q.text) for q in questions),
                         [("gin", OrphanBanks.Q3[1]), ("gq", OrphanBanks.Q1[1])])
        self.assertEqual({c.cid: c.item_count for c in containers}, {"gq": 1, "gin": 1})

    def test_empty_image_attributes_do_not_crash(self):
        for tag in ('<img src="" alt="">', "<img src=''>", "<img>"):
            self.assertEqual(qd.image_token(qd.IMG_RE.search(tag)), " [image: unnamed] ")


class ReviewFixes4(ExportCase):
    """Fixes from the fourth outside review."""

    def test_plain_text_signs_are_not_taken_for_tags(self):
        mat = lambda xml: ET.fromstring(xml)
        self.assertEqual(qd.extract_mattext(mat('<material><mattext texttype="text/plain">'
                                                 'x &lt; 3 and y &gt; 5</mattext></material>')), "x < 3 and y > 5")
        self.assertEqual(qd.strip_html("Is x < 3 and y > 5?"), "Is x < 3 and y > 5?")
        self.assertEqual(qd.strip_html("<p>Is x &lt; 3 <strong>or</strong> x &gt;= 9?</p><!-- c -->"),
                         "Is x < 3 or x >= 9?")

    def test_plain_text_choices_keep_their_signs_through_parsing(self):
        xml = item(MC, "Which is true?", ["A", "B"], ["A"]).replace(
            '<mattext texttype="text/html">A</mattext>', '<mattext texttype="text/plain">x &lt; 3 and y &gt; 5</mattext>')
        (q,), _, _, _ = qd.collect_questions([self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "B", [xml])})])
        self.assertIn("x < 3 and y > 5", q.choices)
        self.assertIn("x < 3 and y > 5", q.correct)

    def test_plain_text_entities_are_decoded(self):
        # a real export had one copy of a choice as text/plain with a trailing &nbsp;
        plain = item(MC, "Which is true?", ["Alpha one", "Beta two"], ["Alpha one"]).replace(
            '<mattext texttype="text/html">Beta two</mattext>',
            '<mattext texttype="text/plain">Beta two&amp;nbsp;</mattext>')
        html_copy = item(MC, "Which is true?", ["Alpha one", "Beta two"], ["Alpha one"])
        questions, _, _, _ = qd.collect_questions([self.export(
            {"non_cc_assessments/gb.xml.qti": bank("gb", "B", [plain, html_copy])})])
        self.assertEqual(questions[0].choices, ["Alpha one", "Beta two"])
        self.assertEqual(len(qd.build_variants(questions)), 1)

    def test_every_stem_block_and_nested_text_is_read(self):
        xml = item(MC, "First paragraph.", ["A one", "B two"], ["A one"]).replace(
            "<presentation>", "<presentation><material><mattext>Read <b>this</b> case.</mattext></material>", 1)
        (q,), _, _, _ = qd.collect_questions([self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "B", [xml])})])
        self.assertEqual(q.text, "Read this case. First paragraph.")
        self.assertEqual(q.choices, ["A one", "B two"])

    def test_quiz_file_found_after_a_long_comment(self):
        xml = bank("gb", "Pool", [item(MC, STEM, CHOICES, [CHOICES[0]])]).replace(
            "<questestinterop", "<!-- " + "x" * 5000 + " -->\n<questestinterop", 1)
        questions = qd.collect_questions([self.export({"non_cc_assessments/gb.xml.qti": xml})])[0]
        self.assertEqual([q.text for q in questions], [STEM])

    def test_uppercase_extensions_are_found(self):
        os.makedirs(os.path.join(self.dir, "x"))
        with open(os.path.join(self.dir, "x", "BANK.QTI"), "w", encoding="utf-8") as f:
            f.write(bank("gb", "B", [item(*OrphanBanks.Q1)]))
        with open(os.path.join(self.dir, "x", "ASSESSMENT_META.XML"), "w", encoding="utf-8") as f:
            f.write("<quiz/>")
        self.assertEqual([os.path.basename(p) for p in qd.find_quiz_xml_files(os.path.join(self.dir, "x"))],
                         ["BANK.QTI"])

    def test_web_images_are_told_apart_by_site_and_path(self):
        qd.IMAGES = qd.ImageStore()
        try:
            name = lambda src: qd.image_token(qd.IMG_RE.search(f'<img src="{src}">'))
            a = name("https://example-a.org/assets/graph.png")
            self.assertNotEqual(a, name("https://example-b.org/charts/graph.png"))
            self.assertEqual(a, name("https://EXAMPLE-A.org/assets/graph.png?verifier=xyz"))
            self.assertEqual(name("https://x.instructure.com/courses/1/files/456/preview?verifier=a"),
                             name("https://x.instructure.com/courses/1/files/456/download"))
        finally:
            qd.IMAGES = None

    def test_min_container_must_be_whole(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "B", [item(*OrphanBanks.Q1)])})
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "--min-container", "2.9"]), \
                contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            qd.main()

    def test_unreadable_pick_count_is_reported(self):
        xml = quiz("gq", "Exam", [], draws=[("gb", 3)]).replace("<selection_number>3<", "<selection_number>three<")
        err = io.StringIO()
        with contextlib.redirect_stderr(err):
            _, _, _, containers = qd.collect_questions([self.export({"non_cc_assessments/gq.xml.qti": xml})])
        ((bank_id, count),) = containers[0].draws
        self.assertEqual((bank_id, type(count), count.text), ("gb", qd.UnreadableCount, "three"))
        self.assertIn("unreadable pick count ('three')", err.getvalue())
        self.assertEqual(qd.describe_contents(containers[0]), "draws an unreadable number from 1 bank")

    def test_unreadable_counts_never_make_quizzes_identical(self):
        files = {f"non_cc_assessments/{q}.xml.qti": quiz(q, title, [], draws=[("gb", 3)]).replace(
                     "<selection_number>3<", "<selection_number>three<")
                 for q, title in (("gq1", "Exam one"), ("gq2", "Exam two"))}
        files["non_cc_assessments/gb.xml.qti"] = bank("gb", "Pool", [item(MC, STEM, CHOICES, [CHOICES[0]])])
        with contextlib.redirect_stderr(io.StringIO()):
            questions, _, _, containers = qd.collect_questions([self.export(files)])
        variants = qd.build_variants(questions)
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        identical, _ = qd.find_duplicate_containers(variants, roots, containers=containers)
        self.assertEqual(identical, [])
        # with a readable count the same two quizzes are identical
        files = {f"non_cc_assessments/{q}.xml.qti": quiz(q, title, [], draws=[("gb", 3)])
                 for q, title in (("gq1", "Exam one"), ("gq2", "Exam two"))}
        files["non_cc_assessments/gb.xml.qti"] = bank("gb", "Pool", [item(MC, STEM, CHOICES, [CHOICES[0]])])
        questions, _, _, containers = qd.collect_questions([self.export(files)])
        variants = qd.build_variants(questions)
        _, roots, _ = qd.find_version_clusters(variants, 85.0)
        identical, _ = qd.find_duplicate_containers(variants, roots, containers=containers)
        self.assertEqual(len(identical), 1)

    def test_outputs_may_not_overwrite_inputs_or_each_other(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Pool", [item(MC, STEM, CHOICES, [CHOICES[0]])])})
        for extra, message in ((["-o", path], "is one of the exports"),
                               (["-o", os.path.join(self.dir, "r.html"), "--json", os.path.join(self.dir, "r.html")],
                                "both write to")):
            err = io.StringIO()
            with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, *extra]), \
                    contextlib.redirect_stderr(err), self.assertRaises(SystemExit):
                qd.main()
            self.assertIn(message, err.getvalue())


class Images(ExportCase):
    IMG = '<img src="$IMS-CC-FILEBASE$/{}?canvas_download=1" alt="x"> What is wrong with this slide?'

    def questions(self, files, refs, as_folder=False):
        """A bank whose questions show the images at `refs`; `files` maps
        archive paths to bytes."""
        entries = {"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, self.IMG.format(r), CHOICES, [CHOICES[0]]) for r in refs])}
        path = self.export(entries)
        with zipfile.ZipFile(path, "a") as z:
            for name, data in files.items():
                z.writestr(name, data)
        if as_folder:
            folder = os.path.join(self.dir, "unpacked")
            with zipfile.ZipFile(path) as z:
                z.extractall(folder)
            path = folder
        store = qd.ImageStore()
        return qd.collect_questions([path], store)[0], store

    def test_same_picture_under_two_names_is_one_question(self):
        qs, store = self.questions({"web_resources/a.png": b"PIC-1", "web_resources/Uploaded Media/b copy.png": b"PIC-1"},
                                   ["a.png", "Uploaded%20Media/b%20copy.png"])
        self.assertEqual(len({q.norm_text for q in qs}), 1)
        self.assertIn("[image: a.png]", qs[1].text)
        self.assertEqual((store.by_content, store.by_name), (2, 0))

    def test_different_pictures_sharing_a_name_are_different_questions(self):
        qs, _ = self.questions({"web_resources/one/slide.png": b"PIC-1", "web_resources/two/slide.png": b"PIC-2"},
                               ["one/slide.png", "two/slide.png"])
        self.assertEqual(len({q.norm_text for q in qs}), 2)
        self.assertIn("[image: slide.png (2)]", qs[1].text)

    def test_missing_file_and_web_address_fall_back_to_the_name(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, self.IMG.format("gone.png"), CHOICES, [CHOICES[0]]),
            item(MC, '<img src="https://example.org/pics/web.jpg"> Describe it.', CHOICES, [CHOICES[0]])])})
        store = qd.ImageStore()
        qs = qd.collect_questions([path], store)[0]
        self.assertIn("[image: gone.png]", qs[0].text)
        self.assertIn("[image: web.jpg]", qs[1].text)
        self.assertEqual((store.by_content, store.by_name, store.by_address), (0, 1, 1))

    def test_same_wording_under_different_pictures_is_not_one_question(self):
        qs, _ = self.questions({"web_resources/s1.png": b"PIC-1", "web_resources/s2.png": b"PIC-2",
                                "web_resources/s1 copy.png": b"PIC-1"}, ["s1.png", "s2.png", "s1%20copy.png"])
        variants = qd.build_variants(qs)
        clusters, _, lookalikes = qd.find_version_clusters(variants, 85.0)
        self.assertEqual((len(variants), clusters, lookalikes), (2, [], []))

    def test_choices_that_are_different_pictures_do_not_match(self):
        stem = "Which of the following graphs represents the function?"
        pic = lambda n: f'<img src="https://example.org/graph{n}.png">'
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, stem, [pic(n) for n in (1, 2, 3, 4)], [pic(1)]),
            item(MC, stem, [pic(n) for n in (5, 6, 7, 8)], [pic(5)]),
            item(MC, stem + " Explain.", [pic(n) for n in (4, 3, 2, 1)], [pic(1)])])})
        variants = qd.build_variants(qd.collect_questions([path])[0])
        self.assertEqual(qd.choice_score(variants[0].choices, variants[1].choices), 0.0)
        clusters = qd.find_version_clusters(variants, 85.0)[0]
        self.assertEqual([sorted(v.rep.item_id for v in c) for c in clusters],
                         [sorted([variants[0].rep.item_id, variants[2].rep.item_id])])

    def test_one_picture_fewer_is_still_an_edit(self):
        both = '<img src="https://example.org/a.png"> <img src="https://example.org/b.png"> ' + STEM
        one = '<img src="https://example.org/a.png"> ' + STEM
        other = '<img src="https://example.org/c.png"> ' + STEM
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, both, CHOICES, [CHOICES[0]]), item(MC, one, CHOICES, [CHOICES[0]]),
            item(MC, other, CHOICES, [CHOICES[0]])])})
        variants = qd.build_variants(qd.collect_questions([path])[0])
        self.assertEqual([len(c) for c in qd.find_version_clusters(variants, 85.0)[0]], [2])

    def test_image_paths_in_other_spellings_are_still_read(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, f'<img src="{src}"> Describe it.', CHOICES, [CHOICES[0]])
            for src in ("$IMS-CC-FILEBASE$/Pics/Chart.PNG", "%24IMS-CC-FILEBASE%24/pics/chart.png",
                        "../web_resources/pics/chart.png", "$IMS_CC_FILEBASE$/pics/chart.png")])})
        with zipfile.ZipFile(path, "a") as z:
            z.writestr("web_resources/pics/chart.png", b"PIC-1")
        store = qd.ImageStore()
        qs = qd.collect_questions([path], store)[0]
        self.assertEqual((store.by_content, store.by_name, len({q.norm_text for q in qs})), (4, 0, 1))

    def test_picture_added_to_a_question_is_still_an_edit(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item(MC, STEM, CHOICES, [CHOICES[0]]),
            item(MC, '<img src="https://example.org/voice.jpg"> ' + STEM, CHOICES, [CHOICES[0]])])})
        variants = qd.build_variants(qd.collect_questions([path])[0])
        self.assertEqual([len(c) for c in qd.find_version_clusters(variants, 85.0)[0]], [2])

    def test_images_are_read_from_an_unpacked_folder_too(self):
        qs, store = self.questions({"web_resources/a.png": b"PIC-1", "web_resources/b.png": b"PIC-1"},
                                   ["a.png", "b.png", "../../outside.png"], as_folder=True)
        self.assertEqual(len({q.norm_text for q in qs[:2]}), 1)
        self.assertEqual((store.by_content, store.by_name), (2, 1))  # the path leaving the folder is not read

    def test_same_picture_in_two_exports_matches(self):
        make = lambda sub, name: self.export({
            "non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [item(MC, self.IMG.format(name), CHOICES, [CHOICES[0]])]),
            "web_resources/" + name: "PIC-1"}, sub=sub)
        qs = qd.collect_questions([make("fall", "old.png"), make("spring", "renamed.png")])[0]
        self.assertEqual(len(qd.build_variants(qs)), 1)


class ReviewFixes(ExportCase):
    def test_exclude_type_does_not_make_quizzes_look_empty_or_questions_missing(self):
        essays = [item("essay_question", p, new=True) for p in list(WholeContainers.STEMS.values())[:3]]
        path = self.export({
            "non_cc_assessments/gb.xml.qti": bank("gb", "Pool", essays, new=True),
            "non_cc_assessments/q1.xml.qti": quiz("q1", "Essay quiz", [item("essay_question", "Discuss ethos.")]),
            "non_cc_assessments/q2.xml.qti": quiz("q2", "Exam", [], picks=[("gb", item_id(essays[0]))], new=True),
        })
        out, target = os.path.join(self.dir, "r.html"), os.path.join(self.dir, "r.json")
        argv = ["quiz_dedup.py", path, "-o", out, "--json", target, "--exclude-type", "essay_question"]
        with mock.patch.object(sys, "argv", argv), contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(target, encoding="utf-8") as f:
            data = json.load(f)
        self.assertEqual(data["summary"]["questions"], 0)
        self.assertEqual(data["quizzes_with_nothing_in_export"], [])
        exam = next(c for c in data["quizzes_and_banks"] if c["title"] == "Exam")
        self.assertEqual(exam["picks_not_found"], 0)

    def test_excluded_bank_is_named_and_excluded_picks_do_not_split_identical_quizzes(self):
        prompts = list(WholeContainers.STEMS.values())
        def course(sub):
            essays = [item("essay_question", p, new=True) for p in prompts[:3]]
            mcq = item(MC, STEM, CHOICES, [CHOICES[0]], new=True)
            return self.export({
                "non_cc_assessments/gb.xml.qti": bank("gb", "Essay pool", essays + [mcq], new=True),
                "non_cc_assessments/gh.xml.qti": bank("gh", "Hidden pool", [item("essay_question", prompts[4], new=True)], new=True),
                "non_cc_assessments/q1.xml.qti": quiz("q1", "Exam", [], new=True, draws=[("gh", 1)],
                                                      picks=[("gb", item_id(essays[0])), ("gb", item_id(mcq))]),
            }, sub=sub, name=sub + ".imscc")
        out, target = os.path.join(self.dir, "r.html"), os.path.join(self.dir, "r.json")
        argv = ["quiz_dedup.py", course("fall"), course("spring"), "-o", out, "--json", target,
                "--exclude-type", "essay_question", "--exclude-title", "hidden"]
        with mock.patch.object(sys, "argv", argv), contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(target, encoding="utf-8") as f:
            data = json.load(f)
        with open(out, encoding="utf-8") as f:
            report = f.read()
        self.assertIn([["Exam", "fall.imscc"], ["Exam", "spring.imscc"]],
                      [sorted([c["title"], c["export"]] for c in g) for g in data["identical_quizzes_and_banks"]])
        self.assertIn("Hidden pool <span class='src'>(left out by --exclude-title)</span>", report)
        # the group is reported, but not as certainly the same: part of what they take was not compared
        self.assertIn("<h3>Possibly the same: 2 quizzes with the same contents", report)
        self.assertNotIn("a bank that is not in this export", report)
        exam = next(c for c in data["quizzes_and_banks"] if c["title"] == "Exam")
        self.assertEqual(exam["draws"][0]["bank_title"], "Hidden pool")

    def test_changed_matching_row_is_marked_in_the_key_line(self):
        rows = [("Communication through touch", "Haptics"), ("Use of personal space", "Proxemics"),
                ("Use of time", "Chronemics"), ("Body movement", "Kinesics")]
        swapped = [(rows[0][0], "Proxemics"), (rows[1][0], "Haptics")] + rows[2:]
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item("matching_question", "Match each term.", rows=rows),
            item("matching_question", "Match each term.", rows=rows),
            item("matching_question", "Match each term.", rows=swapped)])})
        out = os.path.join(self.dir, "r.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        self.assertIn("Communication through touch \u2192 <del>Haptics</del> <ins>Proxemics</ins>", report)
        self.assertIn("Use of personal space \u2192 <del>Proxemics</del> <ins>Haptics</ins>", report)
        self.assertIn("; Use of time \u2192 Chronemics", report)
        self.assertNotIn("<del>Communication through touch", report)

    def test_reworded_matching_prompt_is_diffed_in_place_in_the_key_line(self):
        rows = [("Communication through touch", "Haptics"), ("Use of personal space", "Proxemics"),
                ("Use of time", "Chronemics"), ("Body movement", "Kinesics")]
        reworded = rows[:3] + [("Body movements", "Kinesics")]
        variants = qd.build_variants(qd.collect_questions([self.export({
            "non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
                item("matching_question", "Match each term.", rows=rows),
                item("matching_question", "Match each term.", rows=reworded)])})])[0])
        ref, new = (v.rep for v in sorted(variants, key=lambda v: len(v.rep.choices[3])))
        keyline = qd.render_choices(new, ref)
        keyline = keyline[keyline.index("Answer key"):]
        self.assertIn("Body <del>movement</del> <ins>movements</ins> \u2192 Kinesics", keyline)
        self.assertNotIn("<del>Body movement \u2192", keyline)
        self.assertNotIn("<ins>Body movements \u2192", keyline)
        self.assertFalse(qd.key_differs(ref, new))

    def test_matching_key_moved_to_an_unused_lookalike_option(self):
        ref = [("p1", "apple"), ("p2", "banana"), ("p3", "mitosis phase i")]
        moved = ref[:2] + [("p3", "mitosis phase ii")]
        options = {"apple", "banana", "mitosis phase i", "mitosis phase ii"}
        self.assertTrue(qd.pair_key_differs(ref, moved, options, options))
        # the same answers reworded, with no such option on the other side, is not a key change
        reworded = ref[:2] + [("p3", "mitosis phase one")]
        self.assertFalse(qd.pair_key_differs(ref, reworded, {a for _, a in ref}, {a for _, a in reworded}))

    def test_matching_lists_prompts_before_options(self):
        rows = [("Touch", "Haptics"), ("Space", "Proxemics"), ("Time", "Chronemics")]
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item("matching_question", "Match each term.", rows=rows)])})
        (q,) = qd.collect_questions([path])[0]
        self.assertEqual(q.choices, ["Touch", "Space", "Time", "Chronemics", "Haptics", "Proxemics"])

    def test_short_answer_key_keeps_angle_brackets(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
            item("short_answer_question", "State the range.", answers=["x &lt; 3 and y &gt; 5"])])})
        (q,) = qd.collect_questions([path])[0]
        self.assertEqual(q.correct, ["x < 3 and y > 5"])

    def test_dropdown_options_shared_by_two_blanks_are_not_repeats(self):
        q = qd.Question("x", "Bank", "gb", "i1", "multiple_dropdowns_question", "1", "Rates [a] then [b].",
                        ["increases", "decreases", "increases", "decreases"], "bank")
        mc = qd.Question("x", "Bank", "gb", "i2", MC, "1", "Pick one.", ["up", "down", "up"], "bank")
        repeated = qd.find_repeated_choices(qd.build_variants([q, mc]))
        self.assertEqual([v.qtype for v in repeated], [MC])

    def test_only_quiz_files_are_unpacked(self):
        path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [item("essay_question", "Discuss.")]),
                            "web_resources/lecture.mp4": "x" * 1000, "imsmanifest.xml": "<manifest/>"})
        with tempfile.TemporaryDirectory() as work:
            root = qd.load_source(path, work)
            found = sorted(os.path.relpath(os.path.join(d, f), root) for d, _, fs in os.walk(root) for f in fs)
        self.assertEqual(found, ["imsmanifest.xml", os.path.join("non_cc_assessments", "gb.xml.qti")])

    def test_most_used_label_only_when_one_version_leads(self):
        tweaked = CHOICES[:3] + ["The use of time to send messages"]
        def report(extra):
            path = self.export({"non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [
                item(MC, STEM, CHOICES, [CHOICES[0]]), item(MC, STEM, tweaked, [CHOICES[0]])] + extra)})
            out = os.path.join(self.dir, "r.html")
            with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                    contextlib.redirect_stdout(io.StringIO()):
                qd.main()
            with open(out, encoding="utf-8") as f:
                return f.read()
        self.assertNotIn("(most used)</span>", report([]))
        self.assertIn("(most used)</span>", report([item(MC, STEM, CHOICES, [CHOICES[0]])]))


class TitleFilter(ExportCase):
    """The title filter on the HTML report and the per-bank "also stored
    elsewhere" counts. The filter itself runs in the browser; these check
    that every block carries what it reads. Uses the OrphanBanks course."""
    Q1, Q2, Q2_TWEAKED, Q3, Q4, Q5 = (OrphanBanks.Q1, OrphanBanks.Q2, OrphanBanks.Q2_TWEAKED,
                                      OrphanBanks.Q3, OrphanBanks.Q4, OrphanBanks.Q5)
    course, run_main = OrphanBanks.course, OrphanBanks.run_main

    def blocks(self, report):
        """(tag, data-section, titles) for each element the filter can hide."""
        from html.parser import HTMLParser
        found = []

        class Parser(HTMLParser):
            def handle_starttag(self, tag, attrs):
                attrs = dict(attrs)
                if "data-titles" in attrs:
                    found.append((tag, attrs["data-section"], json.loads(attrs["data-titles"])))

        Parser().feed(report)
        return found

    def test_bank_spread_counts_copies_and_versions_elsewhere(self):
        path = self.export(self.course())
        questions, _, _, containers = qd.collect_questions([path])
        variants = qd.build_variants(questions)
        clusters, _, _ = qd.find_version_clusters(variants, 85.0)
        cluster_of = {id(v): n for n, group in enumerate(clusters, 1) for v in group}
        spread = {cid: counts for (_, cid), counts in qd.bank_spread(variants, containers, cluster_of).items()}
        self.assertEqual(spread["ga"], (2, 2, 0))
        self.assertEqual(spread["gc"], (2, 1, 1))  # Q1 as is, the tweaked Q2 as a version of Q2
        self.assertEqual(spread["gd"], (1, 1, 0))
        self.assertEqual(spread["gf"], (2, 1, 0))  # Q5 is also in the Fallacies quiz, Q4 nowhere else
        self.assertEqual(spread["gg"], (0, 0, 0))
        self.assertNotIn("gq1", spread)  # banks only

    def test_every_block_carries_its_titles(self):
        report, data = self.run_main()
        blocks = self.blocks(report)
        rows = [titles for tag, section, titles in blocks if section == "bank-table"]
        self.assertEqual(sorted(t for (t,) in rows),
                         ["empty", "extra", "old final", "old final", "persuasion", "persuasion", "persuasion"])
        sections = {section for _, section, _ in blocks}
        self.assertTrue({"bank-table", "modified", "exact"} <= sections)
        (modified,) = [titles for _, section, titles in blocks if section == "modified"]
        self.assertEqual(modified, ["persuasion"])
        exact = [titles for _, section, titles in blocks if section == "exact"]
        self.assertIn(["extra", "fallacies quiz"], exact)
        # every section has a count to rewrite and, outside the verdict lists, a note for no match
        for section in sections:
            self.assertIn(f"data-count-for='{section}'" if not section.startswith("verdict-")
                          else f'data-count-for="{section}"', report, section)
            if not section.startswith("verdict-"):
                self.assertIn(f"data-empty-for='{section}'", report, section)
        self.assertIn('id="title-filter"', report)
        self.assertIn("<script>", report)
        self.assertIn("Also stored elsewhere", report)
        extra = next(b for b in data["bank_usage_audit"] if b["id"] == "gf")
        self.assertEqual(extra["also_stored_elsewhere"], {"as_is": 1, "modified": 0})

    def test_quiz_row_matches_the_banks_it_pulls_from(self):
        files = self.course()
        files["non_cc_assessments/gq3.xml.qti"] = quiz("gq3", "Unit 4 quiz", [], draws=[("gd", 1)])
        path = self.export(files)
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            blocks = self.blocks(f.read())
        self.assertIn(["old final", "unit 4 quiz"],
                      [titles for _, section, titles in blocks if section == "from-banks"])

    def test_titles_are_escaped_and_normalised(self):
        files = self.course()
        files["non_cc_assessments/gh.xml.qti"] = bank("gh", html.escape('Chapter  "4" <Language> & Words'), [item(*self.Q4)])
        path = self.export(files)
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            blocks = self.blocks(f.read())
        self.assertIn(['chapter "4" <language> & words'],
                      [titles for _, section, titles in blocks if section == "bank-table"])

    def test_no_match_notes_start_hidden(self):
        import re
        report, _ = self.run_main()
        for note in re.findall(r"<p class='empty' data-empty-for='[^']+'[^>]*>", report):
            self.assertTrue(note.endswith(" hidden>"), note)


class QuizCoverageCheck(ExportCase):
    """Quiz coverage: whether each of a quiz's own questions is also in a
    bank as is, only as a modified version, or nowhere in a bank. Uses the
    OrphanBanks course, where bank ga is linked and the rest are not."""
    Q1, Q2, Q2_TWEAKED, Q3, Q4, Q5 = (OrphanBanks.Q1, OrphanBanks.Q2, OrphanBanks.Q2_TWEAKED,
                                      OrphanBanks.Q3, OrphanBanks.Q4, OrphanBanks.Q5)
    Q5_REWORDED = (MC, "A fallacy that attacks the person rather than the argument is called what?",
                   OrphanBanks.Q5[2], OrphanBanks.Q5[3])
    Q6 = (MC, "Which listening style focuses on the speaker's feelings?",
          ["People-oriented", "Action-oriented", "Content-oriented", "Time-oriented"], ["People-oriented"])

    def course(self):
        files = OrphanBanks.course(self)
        files["non_cc_assessments/gq9.xml.qti"] = quiz("gq9", "Old pool", [
            item(*self.Q1), item(*self.Q1), item(*self.Q3), item(*self.Q5_REWORDED), item(*self.Q6)])
        return files

    def report(self):
        path = self.export(self.course())
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "report.json")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--json", target]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f, open(target, encoding="utf-8") as g:
            return f.read(), json.load(g)

    def test_each_question_lands_in_one_list(self):
        path = self.export(self.course())
        questions, _, _, containers = qd.collect_questions([path])
        usage, _, _ = qd.find_bank_usage(questions, containers)
        variants = qd.build_variants(questions)
        clusters, roots, _ = qd.find_version_clusters(variants, 85.0)
        qd.find_duplicate_containers(variants, roots, containers=containers)
        audit = qd.check_orphan_banks(variants, roots, containers, usage)
        cluster_of = {id(v): n for n, group in enumerate(clusters, 1) for v in group}
        cov = {c.quiz.cid: c for c in qd.quiz_coverage(variants, containers, cluster_of, audit)}
        self.assertNotIn("gq1", cov)  # draws from a bank, stores nothing of its own
        old = cov["gq9"]
        stem = lambda i: variants[i].rep.text
        self.assertEqual(old.distinct, 4)
        self.assertEqual([stem(i) for i in old.in_linked], [self.Q1[1]])  # Q1 is in linked bank ga
        self.assertEqual({stem(i): [k[1] for k in keys] for i, keys in old.in_unlinked.items()},
                         {self.Q3[1]: ["gd", "ge"]})
        ((group, keys),) = old.modified.values()
        self.assertEqual(([k[1] for k in keys], stem(next(iter(old.modified)))), (["gf"], self.Q5_REWORDED[1]))
        self.assertIn(group, cluster_of.values())
        self.assertEqual([stem(i) for i in old.missing], [self.Q6[1]])
        # the Fallacies quiz copy of Q5 is only in Extra, which no quiz links to
        self.assertEqual(len(cov["gq2"].in_unlinked), 1)

    def test_report_section_and_json(self):
        report, data = self.report()
        self.assertIn('id="quiz-coverage"', report)
        self.assertIn("href='#quiz-coverage'", report)
        card = self.card(report, "Old pool")
        self.assertIn("5 questions, 4 distinct", card)
        # Q1 is in the linked Persuasion bank, Q3 in an Old final bank; the linked bank is picked first
        self.assertRegex(card, r"<p>2 of 4 questions are in a bank, spread over 2 banks: <strong>Persuasion</strong> "
                               r"<span class='src'>#\d+</span> \(1\) and <strong>Old final</strong> <span class='src'>#\d+</span> \(1\)\. "
                               r"1 question is in a bank only as a reworded version\. 1 question is in no bank\. ")
        # Old pool holds Q1 twice
        self.assertIn("<strong>It holds 5 questions but only 4 different ones, so a question group picking 4 gives "
                      "students 1 question fewer.</strong>", card)
        self.assertIn("Show the 1 reworded question", card)
        self.assertIn("Show the 1 question in no bank", card)
        self.assertIn(html.escape(self.Q6[1]), card)
        self.assertIn("href='#modified-", card)
        old = next(q for q in data["quiz_coverage"] if q["title"] == "Old pool")
        self.assertEqual((old["in_a_linked_bank"], old["only_in_banks_no_quiz_links_to"],
                          old["only_modified_in_a_bank"], old["in_no_bank"]), (1, 1, 1, 1))
        self.assertEqual([(b["id"], b["questions"]) for b in old["covered_by"]], [("ga", 1), ("gd", 1)])
        self.assertIsNone(old["same_questions_as_bank"])
        self.assertEqual(sorted(q["coverage"] for q in old["questions_to_review"]),
                         ["modified", "none", "only_in_unlinked_banks"])
        self.assertIn(["old pool"], [titles for _, section, titles in TitleFilter.blocks(None, report)
                                     if section == "quiz-coverage"])

    def card(self, report, title):
        start = report.index(f"<strong>{title}</strong>")
        return report[start:report.index("</div>", start)]

    def run_files(self, files):
        path = self.export(files)
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "report.json")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--json", target]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f, open(target, encoding="utf-8") as g:
            return f.read(), json.load(g)

    def test_quiz_with_the_same_questions_as_a_bank(self):
        files = self.course()
        files["non_cc_assessments/gq4.xml.qti"] = quiz("gq4", "Final copy", [item(*self.Q3)])
        report, data = self.run_files(files)
        card = self.card(report, "Final copy")
        self.assertIn("Has the same questions as the bank <strong>Old final</strong>", card)
        self.assertIn("replaced with a question group that picks all 1 from that bank", card)
        self.assertNotIn("<details>", card)  # nothing more to show
        entry = next(q for q in data["quiz_coverage"] if q["title"] == "Final copy")
        self.assertEqual(entry["same_questions_as_bank"]["id"], "gd")

    def test_caution_only_when_a_bank_is_marked_safe(self):
        q7 = (MC, "Hearing is a physiological process; listening is what?",
              ["A psychological process", "A reflex", "A habit", "An accident"], ["A psychological process"])
        files = self.course()
        files["non_cc_assessments/gu.xml.qti"] = bank("gu", "Keeper", [item(*self.Q6), item(*q7)])
        files["non_cc_assessments/gq8.xml.qti"] = quiz("gq8", "Listening quiz", [item(*self.Q6)])
        report, _ = self.run_files(files)
        # Keeper still has a question no quiz holds, so it is not called safe: no caution
        listening = self.card(report, "Listening quiz")
        self.assertIn("Every question is in the bank <strong>Keeper</strong>, which holds 2 questions in all.",
                      listening)
        self.assertNotIn("coverage-warn", listening)
        # Old pool's Q3 is only in the two Old final banks, which its copy makes safe to delete
        old_pool = self.card(report, "Old pool")
        self.assertRegex(old_pool, r"<p class='coverage-warn'>Keep <strong>Old final</strong>.* and "
                                   r"<strong>Old final</strong>.* if you change this quiz\. Some of its questions are "
                                   r"in no bank that will stay")

    def test_no_caution_when_another_quiz_also_has_a_copy(self):
        # the Fallacies quiz copy of Q5 is only in Extra; with a second quiz holding the same copy,
        # changing either quiz alone loses nothing
        files = self.course()
        files["non_cc_assessments/gq6.xml.qti"] = quiz("gq6", "Fallacies quiz 2", [item(*self.Q5)])
        report, _ = self.run_files(files)
        card = self.card(report, "Fallacies quiz 2")
        self.assertIn("Every question is in the bank <strong>Extra</strong>", card)
        self.assertNotIn("coverage-warn", card)

    def test_two_quizzes_sharing_the_only_copies_both_hold_the_bank(self):
        # Q3 is only in the two unlinked Old final banks, and in two quizzes. Each bank is "safe"
        # because the quizzes hold copies; finishing both quizzes would lose Q3, so both must say Keep
        # and the banks must wait for both
        files = self.course()
        files["non_cc_assessments/gq4.xml.qti"] = quiz("gq4", "Second pool", [item(*self.Q3), item(*self.Q6)])
        path = self.export(files)
        out, plan_path = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "plan.xlsx")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--plan", plan_path]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        for title in ("Old pool", "Second pool"):
            self.assertIn("coverage-warn'>Keep <strong>Old final</strong>", self.card(report, title))
        from openpyxl import load_workbook
        rows = list(load_workbook(plan_path)["Banks"].iter_rows(values_only=True))
        head = rows[0]
        old_final = [dict(zip(head, r)) for r in rows[1:] if r[head.index("Bank")].startswith("Old final")]
        self.assertEqual([r["Ready to delete"] for r in old_final], ["wait", "wait"])
        self.assertIn("Old pool and Second pool", old_final[0]["Before deleting"])

    def test_same_named_quizzes_are_numbered(self):
        files = self.course()
        files["non_cc_assessments/gq7.xml.qti"] = quiz("gq7", "Old pool", [item(*self.Q6)])
        report, _ = self.run_files(files)
        section = report[report.index('id="quiz-coverage"'):report.index('id="quizzes-and-banks"')]
        heads = [h for h in section.split("<strong>Old pool</strong>")[1:]]
        self.assertEqual(len(heads), 2)
        for h in heads:
            self.assertTrue(h.startswith(" <span class='src'>#"), h[:60])
            self.assertIn("first question: ", h[:600])

    def test_many_banks_are_summed_up(self):
        stems = ["Which organ pumps blood around the body?", "What gas do plants take in from the air?",
                 "Which planet is known as the red planet?", "What is the boiling point of water at sea level?",
                 "Which ocean is the largest on Earth?", "Which metal is liquid at room temperature?"]
        questions = [(MC, stem, ["Alpha", "Beta", "Gamma", "Delta"], ["Alpha"]) for stem in stems]
        files = self.course()
        for k, q in enumerate(questions):
            files[f"non_cc_assessments/gv{k}.xml.qti"] = bank(f"gv{k}", f"Spread {k}", [item(*q)])
        files["non_cc_assessments/gq5.xml.qti"] = quiz("gq5", "Spread quiz", [item(*q) for q in questions])
        report, _ = self.run_files(files)
        card = self.card(report, "Spread quiz")
        self.assertIn("Every question is in a bank, spread over 6 banks:", card)
        self.assertIn("and 1 more bank (1).", card)
        summary = card[card.index("<p>Every"):]
        self.assertEqual(summary[:summary.index("</p>")].count("<strong>Spread "), 5)
        # every bank is still named beside its question
        self.assertIn("Show which banks hold each question", card)
        self.assertEqual(card[card.index("<details>"):].count("Spread "), 6)

    def test_action_plan(self):
        files = self.course()
        files["non_cc_assessments/gq4.xml.qti"] = quiz("gq4", "Final copy", [item(*self.Q3)])
        path = self.export(files)
        out, plan_path = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "plan.xlsx")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--plan", plan_path]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        from openpyxl import load_workbook
        wb = load_workbook(plan_path)
        self.assertEqual(wb.sheetnames, ["How to use", "Quizzes", "Banks", "Questions", "Repeats"])
        rows = lambda name: [dict(zip(next(wb[name].iter_rows(values_only=True)), r))
                             for r in list(wb[name].iter_rows(values_only=True))[1:]]
        quizzes = {r["Quiz"]: r for r in rows("Quizzes")}
        # Final copy has the same question as an Old final bank: swap for a question group (step 1)
        self.assertEqual(quizzes["Final copy"]["Step"], 1)
        self.assertRegex(quizzes["Final copy"]["Bank"], r"^Old final #\d+$")
        self.assertIn("picks all 1 from the bank", quizzes["Final copy"]["What to do"])
        # the Fallacies quiz's question is in Extra, which holds one more: step 2
        self.assertEqual(quizzes["Fallacies quiz"]["Step"], 2)
        self.assertIn("which has 2", quizzes["Fallacies quiz"]["What to do"])
        # Old pool has a question in no bank and one only reworded: step 4, and both are listed
        self.assertEqual(quizzes["Old pool"]["Step"], 4)
        self.assertEqual([r["Step"] for r in rows("Quizzes")], sorted(r["Step"] for r in rows("Quizzes")))
        questions = [(r["Quiz"], r["Where it is"], r["Question"]) for r in rows("Questions")]
        self.assertIn(("Old pool", "In no bank", self.Q6[1]), questions)
        self.assertIn(("Old pool", "Reworded version in Extra", self.Q5_REWORDED[1]), questions)
        # banks: the one Final copy will draw from is kept; its twin, whose question two quizzes hold, can go
        banks = rows("Banks")
        drawn = quizzes["Final copy"]["Bank"]
        self.assertEqual(next(r for r in banks if r["Bank"] == drawn)["Ready to delete"], "no")
        twin = next(r for r in banks if r["Bank"].startswith("Old final") and r["Bank"] != drawn)
        self.assertEqual(twin["Ready to delete"], "yes")
        self.assertIn("Final copy", twin["Why it is safe"])

    def test_plan_holds_a_bank_only_this_quiz_vouches_for(self):
        files = self.course()  # Old pool alone holds Q3, the only question of both Old final banks
        path = self.export(files)
        questions, _, _, containers = qd.collect_questions([path])
        usage, _, _ = qd.find_bank_usage(questions, containers)
        variants = qd.build_variants(questions)
        clusters, roots, _ = qd.find_version_clusters(variants, 85.0)
        qd.find_duplicate_containers(variants, roots, containers=containers)
        audit = qd.check_orphan_banks(variants, roots, containers, usage)
        cluster_of = {id(v): n for n, group in enumerate(clusters, 1) for v in group}
        coverage = qd.quiz_coverage(variants, containers, cluster_of, audit)
        sizes = {k: held for k, (held, _, _) in qd.bank_spread(variants, [b.bank for b in audit], cluster_of).items()}
        plan = qd.build_plan(variants, coverage, audit, sizes)
        old_final = [r for r in plan["banks"] if r["bank"].startswith("Old final")]
        self.assertEqual([r["ready"] for r in old_final], ["wait", "wait"])
        self.assertIn("Wait: some of its questions are otherwise only in Old pool", old_final[0]["before"])
        self.assertIn("Keep Old final", next(r for r in plan["quizzes"] if r["quiz"] == "Old pool")["keep"])

    def test_a_bank_with_left_out_questions_is_not_the_same(self):
        # Pool has Q1 and Q3 plus an essay; the quiz has Q1 and Q3. With essays left out of the comparison
        # the bank looks like the quiz, but a question group drawing 2 from it could pick the essay
        files = {
            "non_cc_assessments/gp.xml.qti": bank("gp", "Pool", [item(*self.Q1), item(*self.Q3),
                                                                 item("essay_question", "Reflect on ethos.")]),
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Short quiz", [item(*self.Q1), item(*self.Q3)]),
        }
        path = self.export(files)
        out, plan_path = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "plan.xlsx")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--plan", plan_path,
                                             "--exclude-type", "essay_question"]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            card = self.card(f.read(), "Short quiz")
        self.assertIn("Every question is in the bank <strong>Pool</strong>, which holds 3 questions in all.", card)
        from openpyxl import load_workbook
        rows = list(load_workbook(plan_path)["Quizzes"].iter_rows(values_only=True))
        self.assertEqual(rows[1][0], 2)

    def test_questions_left_out_of_the_comparison_are_kept_by_the_plan(self):
        # the quiz has Q1 and Q3, both in Pool, plus an essay; with essays left out the quiz looks like
        # it can be swapped for Pool, but deleting its questions would lose the essay
        files = {
            "non_cc_assessments/gp.xml.qti": bank("gp", "Pool", [item(*self.Q1), item(*self.Q3)]),
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Mixed quiz", [item(*self.Q1), item(*self.Q3),
                                                                       item("essay_question", "Reflect on ethos.")]),
        }
        path = self.export(files)
        out, plan_path = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "plan.xlsx")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--plan", plan_path,
                                             "--exclude-type", "essay_question"]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        from openpyxl import load_workbook
        rows = list(load_workbook(plan_path)["Quizzes"].iter_rows(values_only=True))
        head = rows[0]
        row = dict(zip(head, rows[1]))
        self.assertEqual(row["Step"], 1)
        self.assertTrue(row["What to do"].startswith("Check first. Delete the quiz's questions"), row["What to do"])
        self.assertIn("It also has 1 question left out of the comparison (--exclude-type); it is not covered here, "
                      "so keep it when you delete the others.", row["What to do"])
        with open(out, encoding="utf-8") as f:
            self.assertIn("It also has 1 question left out of the comparison", self.card(f.read(), "Mixed quiz"))

    def test_plan_warns_about_only_published(self):
        path = self.export(self.course())
        plan_path = os.path.join(self.dir, "plan.xlsx")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", os.path.join(self.dir, "r.html"),
                                             "--plan", plan_path, "--only-published"]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        from openpyxl import load_workbook
        text = " ".join(r[0] or "" for r in load_workbook(plan_path)["How to use"].iter_rows(values_only=True))
        self.assertIn("Made with --only-published", text)
        self.assertIn("A quiz in another course can also draw from a bank", text)

    def test_plan_cells_are_never_formulas(self):
        files = self.course()
        files["non_cc_assessments/gq3.xml.qti"] = quiz("gq3", "=SUM(A1:A9)", [
            item(MC, "=1+1 is what?", ["=2", "3", "4", "5"], ["=2"])])
        path = self.export(files)
        out, plan_path = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "plan.xlsx")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--plan", plan_path]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        from openpyxl import load_workbook
        wb = load_workbook(plan_path)
        cells = [c for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                 and c.value.startswith("=")]
        self.assertTrue(cells)
        self.assertEqual({c.data_type for c in cells}, {"s"})
        self.assertIn("=SUM(A1:A9)", {c.value for c in cells})

    def test_keep_ignores_quizzes_in_other_exports(self):
        # Old pool alone holds Q3, the only question of both Old final banks. A quiz in another
        # export with the same question must not lift the Keep note: it cannot stand in for this course
        here = self.export(self.course())
        other = self.export({"non_cc_assessments/gz.xml.qti": quiz("gz", "Other course", [item(*self.Q3)])},
                            name="other.imscc")
        out, plan_path = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "plan.xlsx")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", here, other, "-o", out, "--plan", plan_path]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        from openpyxl import load_workbook
        rows = list(load_workbook(plan_path)["Quizzes"].iter_rows(values_only=True))
        head = rows[0]
        old_pool = next(dict(zip(head, r)) for r in rows[1:] if r[head.index("Quiz")].startswith("Old pool"))
        self.assertIn("Keep Old final", old_pool["Keep"])

    def test_bank_table_shows_distinct_count_when_it_differs(self):
        files = self.course()
        files["non_cc_assessments/gw.xml.qti"] = bank("gw", "Repeats", [item(*self.Q6), item(*self.Q6)])
        path = self.export(files)
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            report = f.read()
        row = report[report.index("<strong>Repeats</strong>"):]
        row = row[:row.index("</tr>")]
        self.assertIn("<td>2 <span class='src'>(1 distinct compared)</span></td>", row)
        persuasion = report[report.index("<strong>Persuasion</strong>"):]
        self.assertIn("<td>2</td>", persuasion[:persuasion.index("</tr>")])

    def test_canvas_error_type_is_explained(self):
        self.assertIn("could not export", qd.type_label("Error"))
        self.assertEqual(qd.type_label(MC), MC)


class Comparison(ExportCase):
    """Side-by-side comparison of two quizzes or banks: in the report under
    each "mostly contained" match, and on its own with --compare."""
    Q1, Q2, Q2_TWEAKED, Q3, Q4, Q5 = (OrphanBanks.Q1, OrphanBanks.Q2, OrphanBanks.Q2_TWEAKED,
                                      OrphanBanks.Q3, OrphanBanks.Q4, OrphanBanks.Q5)
    Q6 = QuizCoverageCheck.Q6

    def course(self):
        a = [self.Q1, self.Q2, self.Q3, self.Q6, self.Q4]
        b = [self.Q1, self.Q2_TWEAKED, self.Q3, self.Q6, self.Q5]
        return {
            "non_cc_assessments/ga.xml.qti": bank("ga", "Pool A", [item(*q) for q in a]),
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Quiz B", [item(*q) for q in b]),
            "non_cc_assessments/gx.xml.qti": quiz("gx", "Twin", [item(*self.Q4)]),
            "non_cc_assessments/gy.xml.qti": quiz("gy", "Twin", [item(*self.Q5)]),
        }

    def run_main(self, *extra):
        path = self.export(self.course())
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, *extra]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        return out

    def test_compare_containers_sorts_every_question(self):
        path = self.export(self.course())
        questions, _, _, containers = qd.collect_questions([path])
        variants = qd.build_variants(questions)
        clusters, roots, _ = qd.find_version_clusters(variants, 85.0)
        qd.find_duplicate_containers(variants, roots, containers=containers)
        cluster_of = {id(v): n for n, group in enumerate(clusters, 1) for v in group}
        a = next(c for c in containers if c.cid == "ga")
        b = next(c for c in containers if c.cid == "gq")
        same, reworded, only_a, only_b = qd.compare_containers(a, b, variants, cluster_of)
        stem = lambda i: variants[i].rep.text
        self.assertEqual(sorted(stem(i) for i in same), sorted([self.Q1[1], self.Q3[1], self.Q6[1]]))
        ((i, j),) = reworded
        self.assertEqual((stem(i), variants[i].rep.choices), (self.Q2[1], self.Q2[2]))
        self.assertEqual(variants[j].rep.choices, self.Q2_TWEAKED[2])
        self.assertEqual([stem(i) for i in only_a], [self.Q4[1]])
        self.assertEqual([stem(i) for i in only_b], [self.Q5[1]])

    def test_report_has_a_side_by_side_under_mostly_contained(self):
        with open(self.run_main(), encoding="utf-8") as f:
            report = f.read()
        section = report[report.index('id="mostly-contained"'):report.index('id="modified"')]
        self.assertIn("Compare side by side", section)
        self.assertIn("<strong>3</strong> the same &middot; <strong>1</strong> reworded &middot; "
                      "<strong>1</strong> only in A &middot; <strong>1</strong> only in B", section)
        self.assertIn("class='compare-grid'", section)
        self.assertIn("The use of time to <del>communicate</del> <ins>send messages</ins>", section)

    def test_long_lists_are_counted_in_the_report(self):
        with mock.patch.object(qd, "COMPARE_LIMIT", 0):
            with open(self.run_main(), encoding="utf-8") as f:
                report = f.read()
        section = report[report.index('id="mostly-contained"'):report.index('id="modified"')]
        self.assertIn("1 reworded question. Run the script with --compare to see them all.", section)
        self.assertNotIn("class='qcard'", section)

    def test_compare_option_writes_a_page(self):
        page = os.path.join(self.dir, "cmp.html")
        out = self.run_main("--compare", "pool a", "quiz", "--compare-output", page)
        self.assertFalse(os.path.exists(out))  # compare mode writes only the comparison
        with open(page, encoding="utf-8") as f:
            html_text = f.read()
        self.assertIn("A is [bank] <strong>Pool A</strong> (5 questions); B is [quiz] <strong>Quiz B</strong>",
                      html_text)
        self.assertIn(".compare-grid", html_text)

    def test_compare_option_names_ambiguous_titles(self):
        with self.assertRaises(SystemExit) as stop:
            self.run_main("--compare", "Twin", "Pool A")
        self.assertIn("matches more than one", str(stop.exception))
        self.assertIn("gx", str(stop.exception))
        with self.assertRaises(SystemExit) as stop:
            self.run_main("--compare", "nothing like this", "Pool A")
        self.assertIn("no quiz or bank matches", str(stop.exception))
        page = os.path.join(self.dir, "cmp.html")
        self.run_main("--compare", "gx", "gy", "--compare-output", page)
        self.assertTrue(os.path.exists(page))


class BankExport(ExportCase):
    """--export-banks: a Canvas course export package with one bank per
    Classic quiz, holding the quiz's questions that are in no bank as is."""
    Q1, Q3, Q5 = OrphanBanks.Q1, OrphanBanks.Q3, OrphanBanks.Q5
    Q5_REWORDED, Q6 = QuizCoverageCheck.Q5_REWORDED, QuizCoverageCheck.Q6

    def course(self):
        broken = item(MC, "Drag the steps into order", CHOICES, [CHOICES[0]]).replace(
            "<fieldentry>multiple_choice_question</fieldentry>", "<fieldentry>Error</fieldentry>")
        return {
            "non_cc_assessments/gb.xml.qti": bank("gb", "Bank", [item(*self.Q1), item(*self.Q5)]),
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Old quiz", [
                item(*self.Q1), item(*self.Q3), item(*self.Q3), item(*self.Q5_REWORDED), item(*self.Q6), broken],
                bank_refs=[]),
            "non_cc_assessments/gn.xml.qti": quiz("gn", "New one", [item(*self.Q6, new=True)], new=True),
            "non_cc_assessments/gd.xml.qti": quiz("gd", "Done quiz", [item(*self.Q1)]),
        }

    def run_export(self):
        path = self.export(self.course())
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "banks.imscc")
        printed = io.StringIO()
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--export-banks", target]), \
                contextlib.redirect_stdout(printed):
            qd.main()
        return path, target, printed.getvalue()

    def test_package_holds_the_questions_not_in_a_bank(self):
        path, target, printed = self.run_export()
        self.assertIn("Bank export written to", printed)
        self.assertIn("1 bank, 3 questions", printed)
        self.assertIn("1 New Quiz with questions in no bank left out", printed)
        self.assertIn("1 question Canvas could not export (type Error) left out", printed)
        with zipfile.ZipFile(target) as z:
            names = z.namelist()
            self.assertIn("imsmanifest.xml", names)
            self.assertIn("course_settings/canvas_export.txt", names)
            manifest = z.read("imsmanifest.xml").decode()
            (bank_file,) = [n for n in names if n.startswith("non_cc_assessments/")]
            self.assertIn(f'href="{bank_file}"', manifest)
            # how Canvas recognises its own export package (lib/canvas/migration/package_identifier.rb):
            # a manifest resource whose href is the flag file, under the Common Cartridge schema
            doc = ET.fromstring(manifest)
            ns = {"m": qd.CC_NS}
            self.assertEqual(doc.findtext("m:metadata/m:schema", namespaces=ns), "IMS Common Cartridge")
            self.assertIsNotNone(doc.find("m:resources/m:resource[@href='course_settings/canvas_export.txt']", ns))
            xml = z.read(bank_file).decode()
        self.assertIn("<fieldentry>Old quiz (questions not in a bank)</fieldentry>", xml)
        self.assertNotIn("assessment_question_identifierref", xml)
        # read back by the script itself: Q3 once, Q6, and the quiz's reworded Q5; Q1 is already in a bank
        questions, _, _, containers = qd.collect_questions([target])
        self.assertEqual([c.title for c in containers if c.kind == "bank"], ["Old quiz (questions not in a bank)"])
        self.assertEqual(sorted(q.text for q in questions), sorted([self.Q3[1], self.Q5_REWORDED[1], self.Q6[1]]))
        self.assertEqual({q.source_kind for q in questions}, {"bank"})
        q6 = next(q for q in questions if q.text == self.Q6[1])
        self.assertEqual((q6.choices, q6.correct, q6.question_type), (self.Q6[2], self.Q6[3], MC))
        originals = {q.item_id for q in qd.collect_questions([path])[0]}
        self.assertFalse({q.item_id for q in questions} & originals)  # fresh idents

    def test_pictures_go_with_their_questions(self):
        stem = '<img src="$IMS-CC-FILEBASE$/Uploaded%20Media/slide.png?canvas_download=1" alt="x"> What is wrong here?'
        path = self.export({"non_cc_assessments/gq.xml.qti": quiz("gq", "Slides quiz", [
            item(MC, stem, CHOICES, [CHOICES[0]])])})
        with zipfile.ZipFile(path, "a") as z:
            z.writestr("web_resources/Uploaded Media/slide.png", b"PNGDATA")
            z.writestr("imsmanifest.xml", '<manifest xmlns="http://www.imsglobal.org/xsd/imsccv1p1/imscp_v1p1">'
                       '<resources><resource identifier="gpic" type="webcontent" '
                       'href="web_resources/Uploaded Media/slide.png"/></resources></manifest>')
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "banks.imscc")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--export-banks", target]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with zipfile.ZipFile(target) as z:
            self.assertEqual(z.read("web_resources/Uploaded Media/slide.png"), b"PNGDATA")
            manifest = z.read("imsmanifest.xml").decode()
        # the picture keeps its export identity, so it maps to the course's existing file
        self.assertIn('identifier="gpic" href="web_resources/Uploaded Media/slide.png"', manifest)

    def test_idents_do_not_depend_on_the_export_file_name(self):
        course = self.course()  # built once: the same Canvas ids, exported under two file names

        def bank_files(name):
            path = self.export(course, name=name)
            target = os.path.join(self.dir, name + ".banks.imscc")
            with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", os.path.join(self.dir, "r.html"),
                                                 "--export-banks", target]), \
                    contextlib.redirect_stdout(io.StringIO()):
                qd.main()
            with zipfile.ZipFile(target) as z:
                return {n: z.read(n) for n in z.namelist() if n.startswith("non_cc_assessments/")}
        self.assertEqual(bank_files("fall-export.imscc"), bank_files("fall-export-2.imscc"))

    def test_pictures_in_other_spellings_keep_the_archive_name(self):
        stem = '<img src="%24IMS-CC-FILEBASE%24/Pics/Chart.PNG" alt="x"> <img src="../web_resources/pics/two.png"> Hm?'
        path = self.export({"non_cc_assessments/gq.xml.qti": quiz("gq", "Charts quiz", [
            item(MC, stem, CHOICES, [CHOICES[0]])])})
        with zipfile.ZipFile(path, "a") as z:
            z.writestr("web_resources/pics/chart.png", b"PIC-1")
            z.writestr("web_resources/pics/two.png", b"PIC-2")
            z.writestr("imsmanifest.xml", '<manifest xmlns="http://www.imsglobal.org/xsd/imsccv1p1/imscp_v1p1">'
                       '<resources><resource identifier="gchart" type="webcontent" '
                       'href="web_resources/pics/chart.png"/></resources></manifest>')
        target = os.path.join(self.dir, "banks.imscc")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", os.path.join(self.dir, "r.html"),
                                             "--export-banks", target]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with zipfile.ZipFile(target) as z:
            self.assertEqual(z.read("web_resources/pics/chart.png"), b"PIC-1")
            self.assertEqual(z.read("web_resources/pics/two.png"), b"PIC-2")
            manifest = z.read("imsmanifest.xml").decode()
        self.assertIn('identifier="gchart" href="web_resources/pics/chart.png"', manifest)

    def test_pictures_in_xml_markup_and_manifest_case_are_matched(self):
        # an <img> written as an XML element inside the mattext, and a manifest href whose case
        # differs from the archive member
        import re
        xml = re.sub(r'<mattext texttype="text/html">[^<]*PLACEHOLDER[^<]*</mattext>',
                     '<mattext texttype="text/html">Look: <img src="$IMS-CC-FILEBASE$/Pics/Map.PNG"/></mattext>',
                     quiz("gq", "Markup quiz", [item(MC, "PLACEHOLDER", CHOICES, [CHOICES[0]])]))
        self.assertIn('<img src="$IMS-CC-FILEBASE$/Pics/Map.PNG"/>', xml)
        path = self.export({"non_cc_assessments/gq.xml.qti": xml})
        with zipfile.ZipFile(path, "a") as z:
            z.writestr("web_resources/Pics/Map.PNG", b"MAP")
            z.writestr("imsmanifest.xml", '<manifest xmlns="http://www.imsglobal.org/xsd/imsccv1p1/imscp_v1p1">'
                       '<resources><resource identifier="gmap" type="webcontent" '
                       'href="web_resources/pics/map.png"/></resources></manifest>')
        target = os.path.join(self.dir, "banks.imscc")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", os.path.join(self.dir, "r.html"),
                                             "--export-banks", target]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with zipfile.ZipFile(target) as z:
            self.assertEqual(z.read("web_resources/Pics/Map.PNG"), b"MAP")
            self.assertIn('identifier="gmap" href="web_resources/Pics/Map.PNG"', z.read("imsmanifest.xml").decode())

    def test_same_titled_quizzes_get_their_report_number(self):
        files = self.course()
        files["non_cc_assessments/gq2.xml.qti"] = quiz("gq2", "Old quiz", [item(*self.Q6)])
        path = self.export(files)
        target = os.path.join(self.dir, "banks.imscc")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", os.path.join(self.dir, "r.html"),
                                             "--export-banks", target]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        titles = sorted(c.title for c in qd.collect_questions([target])[3] if c.kind == "bank")
        self.assertEqual(len(titles), 2)
        for t in titles:
            self.assertRegex(t, r"^Old quiz #\d+ \(questions not in a bank\)$")

    def test_other_outputs_are_refused_in_compare_mode(self):
        path = self.export(self.course())
        err = io.StringIO()
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "--compare", "Bank", "Done quiz",
                                             "--plan", os.path.join(self.dir, "p.xlsx")]), \
                contextlib.redirect_stderr(err), self.assertRaises(SystemExit):
            qd.main()
        self.assertIn("--compare writes only the comparison page; run --plan without --compare", err.getvalue())

    def test_exact_case_wins_over_a_case_twin(self):
        path = os.path.join(self.dir, "twins.zip")
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("web_resources/Image.png", b"UPPER")
            z.writestr("web_resources/image.png", b"LOWER")
        files = qd.ExportFiles(path)
        self.assertEqual((files.read("web_resources/image.png"), files.read("web_resources/Image.png")),
                         (b"LOWER", b"UPPER"))
        self.assertEqual(files.canonical("web_resources/IMAGE.PNG").lower(), "web_resources/image.png")
        files.close()
        store = qd.ImageStore()
        handle = store.open(path)
        self.assertEqual((store.read("web_resources/image.png"), store.read("web_resources/Image.png")),
                         (b"LOWER", b"UPPER"))
        handle.close()

    def test_two_exports_with_different_pictures_at_one_path_keep_their_own(self):
        stem = '<img src="$IMS-CC-FILEBASE$/slide.png"> What is wrong here?'
        paths = []
        for sub, pic, other in (("fall", b"FALL", CHOICES[0]), ("spring", b"SPRING", CHOICES[1])):
            path = self.export({"non_cc_assessments/gq.xml.qti": quiz("gq" + sub, "Slides", [
                item(MC, stem, CHOICES, [other])])}, sub=sub, name=sub + ".imscc")
            with zipfile.ZipFile(path, "a") as z:
                z.writestr("web_resources/slide.png", pic)
            paths.append(path)
        printed = io.StringIO()
        target = os.path.join(self.dir, "banks.imscc")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", *paths, "-o", os.path.join(self.dir, "r.html"),
                                             "--export-banks", target]), \
                contextlib.redirect_stdout(printed):
            qd.main()
        self.assertIn("1 picture from different exports shared a path", printed.getvalue())
        # each bank's question shows its own export's picture
        store = qd.ImageStore()
        questions = qd.collect_questions([target], store)[0]
        self.assertEqual(store.by_content, 2)
        with zipfile.ZipFile(target) as z:
            pictures = {n: z.read(n) for n in z.namelist() if n.endswith("slide.png")}
            manifest = z.read("imsmanifest.xml").decode()
        self.assertEqual(sorted(pictures.values()), [b"FALL", b"SPRING"])
        for name in pictures:
            self.assertIn(f'href="{name}"', manifest)
        by_answer = {q.correct[0]: q.text for q in questions}
        self.assertNotEqual(by_answer[CHOICES[0]], by_answer[CHOICES[1]])  # different pictures, different stems

    def test_quiz_file_with_a_long_header_is_still_exported(self):
        xml = quiz("gq", "Long header quiz", [item(*self.Q6)]).replace(
            "<questestinterop", "<!-- " + "x" * 5000 + " -->\n<questestinterop", 1)
        path = self.export({"non_cc_assessments/gq.xml.qti": xml})
        target = os.path.join(self.dir, "banks.imscc")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", os.path.join(self.dir, "r.html"),
                                             "--export-banks", target, "--max-questions", "50"]), \
                contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            qd.main()
        self.assertEqual([q.text for q in qd.collect_questions([target])[0]], [self.Q6[1]])
        with zipfile.ZipFile(target) as z:
            manifest = z.read("imsmanifest.xml").decode()
        self.assertIn("<!-- Made by quiz_dedup.py with the max-questions option at 50:", manifest)
        ET.fromstring(manifest)  # the comment keeps the manifest well formed

    def test_idents_are_stable_between_runs(self):
        _, target, _ = self.run_export()
        with zipfile.ZipFile(target) as z:
            first = sorted(z.namelist())
        _, target, _ = self.run_export()
        with zipfile.ZipFile(target) as z:
            self.assertEqual(sorted(z.namelist()), first)


class MaxQuestions(ExportCase):
    """--max-questions N leaves out quizzes and banks holding more than N
    questions, so the smaller ones can be compared on their own."""
    QS = [OrphanBanks.Q1, OrphanBanks.Q2, OrphanBanks.Q3, OrphanBanks.Q4, OrphanBanks.Q5]

    def course(self):
        q1, q2, q3 = self.QS[:3]
        return {
            "non_cc_assessments/gbig.xml.qti": bank("gbig", "Publisher bank", [item(*q) for q in self.QS]),
            "non_cc_assessments/gqbig.xml.qti": quiz("gqbig", "Big final", [item(*q) for q in self.QS[:4]]),
            "non_cc_assessments/ga.xml.qti": bank("ga", "Small A", [item(*q1), item(*q2), item(*q3)]),
            "non_cc_assessments/gb.xml.qti": bank("gb", "Small B", [item(*q1), item(*q2), item(*q3)]),
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Unit quiz", [], bank_refs=["ga", "gbig"]),
        }

    def run_main(self, *extra):
        path = self.export(self.course())
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "report.json")
        printed = io.StringIO()
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--json", target, *extra]), \
                contextlib.redirect_stdout(printed):
            qd.main()
        with open(out, encoding="utf-8") as f, open(target, encoding="utf-8") as g:
            return f.read(), json.load(g), printed.getvalue()

    def test_large_quizzes_and_banks_are_left_out(self):
        report, data, printed = self.run_main("--max-questions", "3")
        titles = {c["title"] for c in data["quizzes_and_banks"]}
        self.assertEqual(titles, {"Small A", "Small B", "Unit quiz"})  # the draw-only quiz holds none, so stays
        self.assertEqual(data["summary"]["questions"], 6)
        self.assertEqual(data["settings"]["max_questions"], 3)
        self.assertEqual([(c["title"], c["questions"]) for c in data["left_out_by_max_questions"]],
                         [("Publisher bank", 5), ("Big final", 4)])
        self.assertIn("Left out for holding more than 3 questions (--max-questions): 2 quizzes and banks.", printed)
        self.assertIn("[bank] Publisher bank (5)", printed)
        # the two small banks are still found to be the same
        (group,) = data["identical_quizzes_and_banks"]
        self.assertEqual(sorted(c["title"] for c in group), ["Small A", "Small B"])
        # the report names what was left out, and the quiz still draws from the large bank
        self.assertIn("2 quizzes and banks with more than 3 questions left out (--max-questions)", report)
        self.assertIn("<strong>Publisher bank</strong> <span class='src'>(5 questions)", report)
        self.assertIn("Publisher bank <span class='src'>(left out by --max-questions)</span>", report)
        # Small A is linked; usage was worked out before anything was left out
        self.assertEqual({b["title"]: b["status"] for b in data["bank_usage_audit"]},
                         {"Small A": "linked", "Small B": "unused"})

    def test_without_the_option_nothing_is_left_out(self):
        report, data, _ = self.run_main()
        self.assertEqual(len(data["quizzes_and_banks"]), 5)
        self.assertEqual(data["left_out_by_max_questions"], [])
        self.assertNotIn("--max-questions)</summary>", report)

    def test_limit_must_be_positive(self):
        err = io.StringIO()
        with contextlib.redirect_stderr(err), self.assertRaises(SystemExit):
            self.run_main("--max-questions", "0")
        self.assertIn("--max-questions must be at least 1", err.getvalue())


class InsideDuplicates(ExportCase):
    """Questions one quiz or bank stores more than once, or in more than one version."""
    Q1, Q2, Q2_TWEAKED, Q3 = OrphanBanks.Q1, OrphanBanks.Q2, OrphanBanks.Q2_TWEAKED, OrphanBanks.Q3

    def course(self):
        return {
            # Q1 three times, Q2 in two versions, Q3 once
            "non_cc_assessments/gb.xml.qti": bank("gb", "Messy bank", [
                item(*self.Q1), item(*self.Q1), item(*self.Q1), item(*self.Q2), item(*self.Q2_TWEAKED),
                item(*self.Q3)]),
            # Q1 and Q3 once each: clean, though the questions are also in Messy bank
            "non_cc_assessments/gc.xml.qti": bank("gc", "Clean bank", [item(*self.Q1), item(*self.Q3)]),
            "non_cc_assessments/gq.xml.qti": quiz("gq", "Twice quiz", [item(*self.Q3), item(*self.Q3)]),
        }

    def run_main(self, *extra):
        path = self.export(self.course())
        out, target = os.path.join(self.dir, "report.html"), os.path.join(self.dir, "report.json")
        plan = os.path.join(self.dir, "plan.xlsx")
        printed = io.StringIO()
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out, "--json", target,
                                             "--plan", plan, *extra]), contextlib.redirect_stdout(printed):
            qd.main()
        with open(out, encoding="utf-8") as f, open(target, encoding="utf-8") as g:
            return f.read(), json.load(g), plan, printed.getvalue()

    def test_repeats_and_versions_inside_one_container(self):
        report, data, plan, printed = self.run_main()
        inside = {d["title"]: d for d in data["duplicates_inside"]}
        self.assertEqual(set(inside), {"Messy bank", "Twice quiz"})  # copies across banks do not count
        messy = inside["Messy bank"]
        self.assertEqual([(r["stem"], r["copies"]) for r in messy["stored_more_than_once"]], [(self.Q1[1], 3)])
        (versions,) = messy["in_more_than_one_version"]
        self.assertEqual(versions["stems"], [self.Q2[1], self.Q2_TWEAKED[1]])
        self.assertIsInstance(versions["modified_group"], int)
        self.assertEqual([(r["stem"], r["copies"]) for r in inside["Twice quiz"]["stored_more_than_once"]],
                         [(self.Q3[1], 2)])
        self.assertIn("Duplicates inside one quiz or bank: 2 quizzes and banks, 3 extra copies and 1 question "
                      "in more than one version", printed)
        # the report: a section with a card per container, in the contents and the filter
        section = report[report.index('<h2 id="inside">'):report.index('<h2 id="modified">')]
        self.assertIn("1 question stored more than once (2 extra copies); 1 question stored in more than one "
                      "version.", section)
        self.assertIn("&times;3", section)
        self.assertIn("The use of time to <del>communicate</del> <ins>send messages</ins>", section)
        self.assertIn(f"href='#modified-{versions['modified_group']}'", section)
        self.assertIn("href='#inside'", report)
        self.assertIn('data-section="inside"', section)
        # the plan's Repeats tab: one row per thing to fix
        from openpyxl import load_workbook
        rows = list(load_workbook(plan)["Repeats"].iter_rows(values_only=True))
        self.assertEqual(sorted(r[2] for r in rows[1:]),
                         ["2 versions: keep one wording", "Stored 2 times: delete 1", "Stored 3 times: delete 2"])

    def test_nothing_found(self):
        path = self.export({"non_cc_assessments/gc.xml.qti": bank("gc", "Clean bank", [item(*self.Q1)])})
        out = os.path.join(self.dir, "report.html")
        with mock.patch.object(sys, "argv", ["quiz_dedup.py", path, "-o", out]), \
                contextlib.redirect_stdout(io.StringIO()):
            qd.main()
        with open(out, encoding="utf-8") as f:
            self.assertIn("No quiz or bank holds a question twice or in two versions.", f.read())


class Scoring(unittest.TestCase):
    def test_numbers_keep_their_sign_and_shape(self):
        forms = ["5", "-5", ".5", "1.5", "15", "1-5", "\u22125"]
        norm = [qd.normalize(x) for x in forms]
        self.assertEqual(len(set(norm)), 6)
        self.assertEqual(norm[1], norm[6])  # a typed hyphen and a true minus sign
        self.assertEqual(qd.normalize("audience-analysis"), "audienceanalysis")
        self.assertEqual(qd.normalize("0.5"), qd.normalize(".5"))

    def test_decimals_are_kept_apart(self):
        self.assertNotEqual(qd.normalize("Add 1.5 mL"), qd.normalize("Add 15 mL"))
        self.assertEqual(qd.normalize("Stop. Go."), "stop go")

    def test_word_overlap_of_unrelated_stems_is_zero(self):
        self.assertEqual(qd.word_overlap("", "explain the model", {}), 0.0)
        self.assertEqual(qd.word_overlap("alpha beta", "gamma delta", {}), 0.0)

    def test_choice_pairing_is_one_to_one(self):
        pairs = qd.align_choices(["car", "cars", "truck"], ["car", "truck", "bus"])
        self.assertEqual(len({i for i, _, _ in pairs}), 3)
        self.assertEqual(len({j for _, j, _ in pairs}), 3)
        self.assertLess(qd.choice_score(frozenset({"car", "cars", "truck"}), frozenset({"car", "truck", "bus"})), 80)

    def test_choice_score_does_not_depend_on_set_order(self):
        # "aa" and "ab" both tie on "ac"; which one claims it decides the total
        a, b = ["aa", "ab"], ["ac", "bc"]
        scores = {qd.choice_score(dict.fromkeys(p), dict.fromkeys(q))
                  for p in (a, a[::-1]) for q in (b, b[::-1])}
        self.assertEqual(len(scores), 1)

    def test_word_diff_marks_only_changed_words(self):
        out = qd.word_diff("the quick brown fox", "the slow brown fox")
        self.assertEqual(out, "the <del>quick</del> <ins>slow</ins> brown fox")


if __name__ == "__main__":
    unittest.main()
